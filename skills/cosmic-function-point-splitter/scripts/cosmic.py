#!/usr/bin/env python3
"""COSMIC 功能点拆分 CLI。

    inspect   读送审表模板，打印列布局与合并规则
    init      生成规格文件骨架
    check     校验规格并预估功能点，不写文件
    build     生成拆分表（默认同时跑结构校验，不通过就不落盘）
    audit     核对已有表的数据移动、过程边界与原表计数
    report    输出送审自检清单
    coverage  核对接口/页面清单与规格的双向覆盖
    normalize 把已有表的子过程描述改成当前句式（默认干跑）

结构错误与一致性提示分开：前者阻塞，后者只提示；--strict 把提示也当失败。
"""

from __future__ import annotations

import argparse
import json
import shutil
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import builder  # noqa: E402
import checks  # noqa: E402
import inventory as inventory_module  # noqa: E402  # pyright: ignore[reportMissingImports]
import operation_types as ops  # noqa: E402  # pyright: ignore[reportMissingImports]
import sheet  # noqa: E402
import spec as spec_module  # noqa: E402
import style  # noqa: E402  # pyright: ignore[reportMissingImports]

STARTER_SPEC = {
    "version": 3, "project": "示例项目",
    "measurement": {"purpose": "功能规模度量", "scope": "地址管理服务",
                    "boundary": "用户与地址管理服务之间；持久存储在边界内",
                    "layer": "应用层", "functional_users": ["业务用户"],
                    "conventions": {"描述句式": "简洁动宾式：输入/读取/保存/返回"}},
    "modules": [{"l2": "示例一级模块", "l3": "地址管理"}],
    "processes": [{"id": "FP-001", "l3": "地址管理", "name": "地址新增", "op": "insert",
                   "trigger": "业务用户提交新地址", "user": "业务用户", "status": "unverified",
                   "op_note": "新增前读取已有地址查重，属于同一触发下的完整响应",
                   "movements": [
        {"type": "E", "object": "地址", "group": "地址资料", "attributes": ["地址内容"], "description": "输入地址新增请求"},
        {"type": "R", "object": "地址", "group": "地址资料", "attributes": ["地址内容"], "description": "读取地址信息"},
        {"type": "W", "object": "地址", "group": "地址资料", "attributes": ["地址内容"], "description": "保存地址信息"},
        {"type": "X", "object": "操作结果", "group": "处理结果", "attributes": ["结果状态"], "description": "返回地址新增结果"},
    ]}],
}


def _fail(message: str) -> int:
    print(message, file=sys.stderr)
    return 1


def _path_conflict(paths: dict[str, str | Path | None]) -> str | None:
    """写入前识别输入、输出之间的同路径及链接别名，保留正常同名输出重建。"""
    items = [(label, Path(value)) for label, value in paths.items() if value is not None]
    for index, (left_label, left) in enumerate(items):
        for right_label, right in items[index + 1:]:
            try:
                same = left.resolve() == right.resolve()
                if not same and left.exists() and right.exists():
                    same = left.samefile(right)
            except (OSError, RuntimeError) as exc:
                return f"无法核对输出路径：{exc}"
            if same:
                return f"路径冲突：{left_label}与{right_label}指向同一文件（{left}），未写出文件"
    return None


def command_inspect(args: argparse.Namespace) -> int:
    template = sheet.inspect(args.template, args.sheet)
    print(sheet.describe(template))
    return 0


def command_init(args: argparse.Namespace) -> int:
    target = Path(args.path)
    if target.exists() and not args.force:
        return _fail(f"{target} 已存在；加 --force 覆盖")
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(
        json.dumps(STARTER_SPEC, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(f"已写出规格骨架：{target}")
    print("下一步：改成你的度量范围、功能过程和实际数据移动，然后跑 check。")
    return 0


def _load(args: argparse.Namespace) -> dict:
    spec = spec_module.load(args.spec)
    if args.project:
        spec["project"] = args.project
    return spec


def _print_stats(spec: dict) -> None:
    info = spec_module.stats(spec)
    print("| 二级模块 | 三级模块 | 功能过程 | CFP |")
    print("| --- | --- | --- | --- |")
    for l2, entry in info["modules"].items():
        print(f"| {l2} | {entry['l3']} | {entry['processes']} | {entry['cfp']} |")
    print(f"\n合计 {info['processes']} 个功能过程、{info['cfp']} CFP")


def _coverage(args: argparse.Namespace, spec: dict):
    if not args.inventory:
        return None
    return inventory_module.coverage(spec, inventory_module.load(args.inventory))


def _print_report(hard: list[str], hints: list[str], coverage, args: argparse.Namespace) -> None:
    print()
    if hard:
        print(f"结构问题 {len(hard)} 个：")
        for issue in hard[: args.max_issues]:
            print(f"  - {issue}")
        if len(hard) > args.max_issues:
            print(f"  ……另有 {len(hard) - args.max_issues} 个未显示")
    if hints:
        print(("一致性提示（--strict 时阻塞）" if getattr(args, "strict", False) else "一致性提示（不阻塞）")
              + f" {len(hints)} 个：")
        for hint in hints[: args.max_issues]:
            print(f"  - {hint}")
        if len(hints) > args.max_issues:
            print(f"  ……另有 {len(hints) - args.max_issues} 个未显示")
    if coverage is not None:
        print()
        print(inventory_module.render(coverage))
        if not coverage.ok:
            print("  （清单与规格的差异不自动等于漏拆；确认范围后把接口 id 写进过程的 evidence）")


def _blocked(hard: list[str], hints: list[str], coverage, args: argparse.Namespace) -> bool:
    if hard:
        return True
    if getattr(args, "strict", False) and hints:
        return True
    return bool(getattr(args, "strict", False) and coverage is not None and not coverage.ok)


def _measurement(rows: list[builder.Row], spec: dict) -> list[str]:
    summary = checks.summary(rows)
    lines = [f"候选合计 {summary['cfp']} CFP；已确认标记 {summary['confirmed_cfp']} CFP；"
             f"待核实 {summary['unverified_cfp']} CFP（{summary['unverified_processes']} 个过程）。",
             "已确认仅表示规格作者声明已核实；工具无法验证需求内容。目标区间不参与计算。"]
    for key, value in spec.get("measurement", {}).items():
        if key == "conventions" and isinstance(value, dict):
            for name, text in value.items():
                lines.append(f"- 约定·{name}：{text}")
        elif key == "functional_users":
            lines.append(f"- functional_users：{'、'.join(value)}")
        else:
            lines.append(f"- {key}：{value}")
    return lines


def command_check(args: argparse.Namespace) -> int:
    spec = _load(args)
    _print_stats(spec)
    result = builder.build(spec)
    hard = checks.check(result.rows)
    hints = checks.hints(result.rows, style.load_glossary(args.glossary))
    coverage = _coverage(args, spec)
    _print_report(hard, hints, coverage, args)
    if _blocked(hard, hints, coverage, args):
        return 1
    print("校验通过：结构检查通过，不代表来源已核实或符合全部 COSMIC 规则。")
    _print_measurement_summary(result.rows)
    return 0


def _unconfirmed(rows: list[builder.Row]) -> int:
    return sum(1 for row in rows if row.is_process_head and row.status != "confirmed")


def command_build(args: argparse.Namespace) -> int:
    report_path = args.selfcheck or str(args.output) + ".selfcheck.md"
    conflict = _path_conflict({"规格": args.spec, "模板": args.template,
                               "接口清单": args.inventory,
                               "术语表": args.glossary if isinstance(args.glossary, str) else None,
                               "拆分表": args.output, "来源报告": report_path})
    if conflict:
        return _fail(conflict)
    spec = _load(args)
    result = builder.build(spec, order=args.order)
    hard = checks.check(result.rows)
    hints = checks.hints(result.rows, style.load_glossary(args.glossary))
    coverage = _coverage(args, spec)
    if args.require_confirmed:
        pending = _unconfirmed(result.rows)
        if pending:
            hard.append(f"--require-confirmed：还有 {pending} 个功能过程未标 confirmed")
    if _blocked(hard, hints, coverage, args) and not args.force:
        print(f"校验未通过，{len(hard)} 个结构问题、{len(hints)} 个提示，未写出文件：", file=sys.stderr)
        for issue in hard[: args.max_issues]:
            print(f"  - {issue}", file=sys.stderr)
        for hint in hints[: args.max_issues]:
            print(f"  - 提示：{hint}", file=sys.stderr)
        if coverage is not None and not coverage.ok:
            print(inventory_module.render(coverage), file=sys.stderr)
        print("修好后再跑，或用 --force 强制写出（不推荐）。", file=sys.stderr)
        return 1

    template = sheet.inspect(args.template, args.sheet) if args.template else sheet.Template()
    template.l1_value = args.l1 or spec["project"]

    output = sheet.write(result.rows, args.output, template, user_format=args.user_format)
    summary = checks.summary(result.rows)
    print(f"已写出：{output}")
    print(f"  {summary['processes']} 个功能过程 / {summary['rows']} 行 / {summary['cfp']} CFP")
    print(f"  数据移动分布：{summary['moves']}")
    _print_measurement_summary(result.rows)
    if template.notes:
        for note in template.notes:
            print(f"  模板：{note}")
    if hard:
        print(f"！--force 跳过校验，仍有 {len(hard)} 个结构问题、{len(hints)} 个提示；产物为待修正草稿")
    elif hints:
        print(f"  一致性提示 {len(hints)} 个，已写入自检清单（不阻塞）")
    report_text = selfcheck_report(result.rows, summary, hard, spec, hints, coverage)
    Path(report_path).parent.mkdir(parents=True, exist_ok=True)
    Path(report_path).write_text(report_text, encoding="utf-8")
    print(f"已写出自检及来源清单：{report_path}")
    return 0


def command_report(args: argparse.Namespace) -> int:
    conflict = _path_conflict({"规格": args.spec, "接口清单": args.inventory,
                               "术语表": args.glossary if isinstance(args.glossary, str) else None,
                               "报告": args.output})
    if conflict:
        return _fail(conflict)
    spec = _load(args)
    result = builder.build(spec)
    hard = checks.check(result.rows)
    hints = checks.hints(result.rows, style.load_glossary(args.glossary))
    coverage = _coverage(args, spec)
    summary = checks.summary(result.rows)
    text = selfcheck_report(result.rows, summary, hard, spec, hints, coverage)
    if args.output:
        Path(args.output).parent.mkdir(parents=True, exist_ok=True)
        Path(args.output).write_text(text, encoding="utf-8")
        print(f"已写出自检清单：{args.output}")
    else:
        print(text)
    return 0 if not _blocked(hard, hints, coverage, args) else 1


def command_audit(args: argparse.Namespace) -> int:
    rows = read_existing_table(args.xlsx, args.sheet)
    if not rows:
        return _fail(f"{args.xlsx} 里没读到数据行；用 inspect 先确认列布局是否被识别")
    hard = checks.check(rows)
    hints = checks.hints(rows, style.load_glossary(args.glossary))
    print(f"读到 {len(rows)} 行、{sum(1 for row in rows if row.is_process_head)} 个功能过程")
    _print_measurement_summary(rows)
    print("审计仅检查表内结构与写法；未核实原始需求和计量边界。")
    _print_report(hard, hints, None, args)
    if not hard:
        print("未发现结构问题。")
    return 0 if not _blocked(hard, hints, None, args) else 1


def command_coverage(args: argparse.Namespace) -> int:
    conflict = _path_conflict({"规格": args.spec, "接口清单": args.inventory, "报告": args.output})
    if conflict:
        return _fail(conflict)
    spec = spec_module.load(args.spec)
    inventory = inventory_module.load(args.inventory)
    result = inventory_module.coverage(spec, inventory)
    payload = inventory_module.as_dict(result)
    if args.json:
        text = json.dumps(payload, ensure_ascii=False, indent=2)
    else:
        text = inventory_module.render(result)
    if args.output:
        Path(args.output).parent.mkdir(parents=True, exist_ok=True)
        Path(args.output).write_text(text + "\n", encoding="utf-8")
        print(f"已写出覆盖报告：{args.output}")
    else:
        print(text)
    return 0 if result.ok else 1


def command_normalize(args: argparse.Namespace) -> int:
    source = Path(args.xlsx)
    glossary = style.load_glossary(args.glossary) if args.glossary else None
    changes = normalize_changes(source, args.sheet, glossary)
    if not changes:
        print("没有需要改写的子过程描述。")
        return 0
    print("| 行 | 列 | 原文 | 改写 |")
    print("| --- | --- | --- | --- |")
    for row, label, old, new in changes[: args.max_issues]:
        print(f"| {row} | {label} | {old} | {new} |")
    if len(changes) > args.max_issues:
        print(f"……另有 {len(changes) - args.max_issues} 处未显示")
    if not args.write:
        print(f"\n干跑：{len(changes)} 处待改写，未写入。加 -o 目标文件与 --write 执行。")
        return 0
    if not args.output:
        return _fail("--write 需要同时指定 -o 输出文件；不会覆盖原表")
    conflict = _path_conflict({"原表": source, "输出": args.output,
                               "术语表": args.glossary if isinstance(args.glossary, str) else None})
    if conflict:
        return _fail(conflict)
    written = write_normalized(source, args.output, changes, args.sheet)
    print(f"\n已写出：{written}（{len(changes)} 处改写，其余内容、合并与样式原样保留）")
    return 0


def _desc_cells(source: Path, sheet_name: str | None):
    openpyxl = sheet._openpyxl()
    template = sheet.inspect(source, sheet_name)
    workbook = openpyxl.load_workbook(source)
    return workbook, workbook[template.sheet_name], template.layout


def normalize_changes(source: Path, sheet_name: str | None,
                      glossary: dict | None) -> list[tuple[int, str, str, str]]:
    """只改文字：子过程描述始终参与；给了术语表时也清理数据组与属性。"""
    workbook, worksheet, layout = _desc_cells(source, sheet_name)
    try:
        rows = read_existing_table(str(source), sheet_name)
        changes = []
        for row in rows:
            targets = [("desc", "子过程描述", row.desc)]
            if glossary:
                targets += [("group", "数据组", row.group), ("attrs", "数据属性", row.attrs)]
            for key, label, text in targets:
                column = layout.get(key)
                if not column or not _writable_cell(worksheet, row.excel_row, column):
                    continue
                cell = worksheet.cell(row.excel_row, column)
                if cell.data_type == "f" or not isinstance(cell.value, str):
                    continue
                text = cell.value
                new = style.rewrite_description(text, row.move, row.func) if key == "desc" else text
                if glossary:
                    new, _ = style.replace_english(new, glossary)
                if new != text:
                    changes.append((row.excel_row, label, text, new))
        return changes
    finally:
        workbook.close()


def write_normalized(source: Path, output: str | Path, changes: list[tuple[int, str, str, str]],
                     sheet_name: str | None = None) -> Path:
    """复制原表后只改目标单元格，避免 insert/delete_rows 破坏合并单元格。"""
    target = Path(output)
    for row, label, _old, new in changes:
        sheet.validate_cell_text(new, f"第 {row} 行 {label}")
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, target)
    openpyxl = sheet._openpyxl()
    workbook = openpyxl.load_workbook(target)
    template = sheet.inspect(target, sheet_name)
    worksheet = workbook[template.sheet_name]
    labels = {"子过程描述": "desc", "数据组": "group", "数据属性": "attrs"}
    columns = {label: template.layout[key] for label, key in labels.items() if key in template.layout}
    for row, label, _old, new in changes:
        column = columns.get(label)
        if not column:
            continue
        if not _writable_cell(worksheet, row, column):
            continue
        from openpyxl.cell.cell import Cell
        cell = worksheet.cell(row=row, column=column)
        if not isinstance(cell, Cell):  # 合并区非左上角在运行期也是 MergedCell
            continue
        cell.value = new
        cell.data_type = "s"
    workbook.save(target)
    workbook.close()
    return target


def _writable_cell(worksheet, row: int, column: int) -> bool:
    """合并区内只有左上角可写；其余位置由 openpyxl 的 MergedCell 表示。"""
    for rng in worksheet.merged_cells.ranges:
        if (rng.min_row <= row <= rng.max_row) and (rng.min_col <= column <= rng.max_col):
            return (row, column) == (rng.min_row, rng.min_col)
    return True


def read_existing_table(path: str, sheet_name: str | None = None) -> list[builder.Row]:
    """使用探测选中的表，保留非法行及原始 CFP，按过程标识分组。"""
    openpyxl = sheet._openpyxl()
    template = sheet.inspect(path, sheet_name)
    workbook = openpyxl.load_workbook(path, data_only=False)
    worksheet = workbook[template.sheet_name]
    layout = template.layout
    filled = {}
    for rng in worksheet.merged_cells.ranges:
        value = worksheet.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                filled[(r, c)] = value

    def cell(key, r, raw=False):
        c = layout.get(key)
        if not c:
            return None
        value = worksheet.cell(r, c).value
        return value if raw else filled.get((r, c), value)

    def txt(key, r):
        value = cell(key, r)
        return str(value).strip() if value is not None else ''

    rows = []
    last = {}
    for r in range(template.header_row + 1, worksheet.max_row + 1):
        # 只跳过没有原始业务内容的空行；合并单元格的延伸不能凭空生成行。
        # 缺损过程及未识别的汇总行都保留，不能仅因没有移动字段就静默忽略。
        if not any(cell(k, r, raw=True) not in (None, '')
                   for k in layout if k not in ('l1', 'reuse')):
            continue
        current = {k: txt(k, r) for k in ('l2', 'l3', 'user', 'event', 'func', 'process_id')}
        changed_module = any(current[k] and current[k] != last.get(k) for k in ('l2', 'l3'))
        new_process = (not last
                       or bool(current['process_id'] and current['process_id'] != last.get('process_id'))
                       or bool(current['func'] and current['func'] != last.get('func'))
                       or (changed_module and not current['process_id']))
        values = {k: current[k] or ('' if new_process else last.get(k, ''))
                  for k in ('user', 'event', 'func', 'process_id')}
        values['l2'] = current['l2'] or last.get('l2', '')
        values['l3'] = current['l3'] or (
            '' if current['l2'] and current['l2'] != last.get('l2') else last.get('l3', ''))
        last = values.copy()
        row = builder.Row(
            l2=values['l2'], l3=values['l3'], user=values['user'], event=values['event'],
            func=values['func'], process_id=values['process_id'],
            desc=txt('desc', r), move=txt('move', r).upper(),
            group=txt('group', r), attrs=txt('attrs', r),
            excel_row=r, declared_cfp=cell('cfp', r, raw=True),
            status=txt('status', r) or 'unverified', source=txt('source', r),
            object=txt('object', r), repeat_reason=txt('repeat_reason', r),
            op=txt('op', r), op_note=txt('op_note', r), evidence=txt('evidence', r))
        if not any(cell(k, r) not in (None, '') for k in ('move', 'desc', 'group', 'attrs')):
            row.read_issues.append(f"无法识别为数据移动的非空行 第 {r} 行：需核实是否为缺损过程或汇总行")
        rows.append(row)
    required = ('l2', 'l3', 'user', 'event', 'func', 'desc', 'move', 'group', 'attrs', 'cfp', 'object')
    missing = [sheet.HEADER_LABELS[k] for k in required if k not in layout]
    if rows and missing:
        rows[0].read_issues.append('原表缺少核心列：' + '、'.join(missing))
    for block in checks.process_blocks(rows):
        first = block[0]
        declarations = [row.declared_cfp for row in block if row.declared_cfp not in (None, '')]
        numbers = [v for v in declarations if isinstance(v, (int, float)) and not isinstance(v, bool)]
        valid = len(numbers) == len(declarations) and all(v >= 0 and v == int(v) for v in numbers)
        if not declarations:
            first.read_issues.append(f"原表缺少 CFP 声明：{first.func}，重算值仅供参考")
        elif not valid:
            first.read_issues.append(f"原表 CFP 非法或为未求值公式：{first.func} {declarations!r}")
        first.fp = int(sum(numbers)) if valid and numbers else len(block)
        for index, row in enumerate(block):
            row.seq_index = index
            if index:
                row.fp = 0
            if row.status not in ('confirmed', 'unverified'):
                row.read_issues.append(f"状态非法 第 {row.excel_row} 行：{row.status}")
            if row.status == 'confirmed' and not row.source:
                row.read_issues.append(f"已确认项缺来源 第 {row.excel_row} 行")
    workbook.close()
    return rows


def _print_measurement_summary(rows: list[builder.Row]) -> None:
    info = checks.summary(rows)
    print(f"  已确认标记 {info['confirmed_cfp']} CFP；待核实草稿 {info['unverified_cfp']} CFP"
          f"（{info['unverified_processes']} 个过程）")
    print("  已确认指规格作者的标记，工具未自动验证来源。目标数不参与计量。")


def selfcheck_report(rows: list[builder.Row], summary: dict, issues: list[str], spec: dict | None = None,
                     hints: list[str] | None = None, coverage=None) -> str:
    def safe(value):
        return str(value).replace('|', r'\|').replace('\n', ' ')
    lines = ["# COSMIC 拆分表自检清单", "",
             f"结构检查：{'未通过，产物为待修正草稿' if issues else '通过（不等于标准认证或可送审）'}", ""]
    if spec:
        lines += _measurement(rows, spec) + [""]
    else:
        lines += [f"候选合计 {summary['cfp']} CFP；已确认标记 {summary['confirmed_cfp']} CFP；"
                  f"待核实 {summary['unverified_cfp']} CFP（{summary['unverified_processes']} 个过程）。", ""]
    if coverage is not None:
        lines += ["## 接口与页面覆盖", "", "```", inventory_module.render(coverage), "```", ""]
    lines += ["## 一致性提示", ""]
    lines += [f"- {safe(hint)}" for hint in (hints or [])] or ["- 无。句式与排序的本地惯例未发现偏离。"]
    lines += ["", "## 逐项来源与计量依据", "",
              "| 过程编号 | op | 操作依据 | 功能过程 | 状态 | 触发事件 | 功能用户 | 移动 | 兴趣对象 | 数据组 | 属性 | 来源 | 接口证据 | 重复移动理由 |",
              "| --- | --- | --- | --- | --- | --- | --- | --- | --- | --- | --- | --- | --- | --- |"]
    for r in rows:
        lines.append('| ' + ' | '.join(safe(v) for v in (r.process_id, r.op or '', r.op_note, r.func, r.status, r.event,
                     r.user, r.move, r.object, r.group, r.attrs, r.source or '待核实',
                     r.evidence or '—', r.repeat_reason)) + ' |')
    lines += ["", "## 待处理问题", ""]
    lines += [f"- {safe(issue)}" for issue in issues] or ["- 未发现结构问题；仍需人工核对边界、来源、过程划分和重复移动理由。"]
    return '\n'.join(lines) + '\n'


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="cosmic.py", description="COSMIC 功能点拆分：规格 → 拆分表 + 全表校验"
    )
    sub = parser.add_subparsers(dest="command", required=True)

    inspect_parser = sub.add_parser("inspect", help="读送审表模板，打印列布局与合并规则")
    inspect_parser.add_argument("template")
    inspect_parser.add_argument("--sheet", help="工作表名，默认自动探测")
    inspect_parser.set_defaults(handler=command_inspect)

    init_parser = sub.add_parser("init", help="写出规格文件骨架")
    init_parser.add_argument("path")
    init_parser.add_argument("--force", action="store_true")
    init_parser.set_defaults(handler=command_init)

    common = argparse.ArgumentParser(add_help=False)
    common.add_argument("spec", help="规格 JSON")
    common.add_argument("--project", help="覆盖规格里的项目名（一级模块取值）")
    common.add_argument("--max-issues", type=int, default=25)
    common.add_argument("--strict", action="store_true", help="把一致性提示与覆盖缺口也当作失败")
    common.add_argument("--inventory", help="接口与页面清单 JSON，用于覆盖核对")
    common.add_argument("--glossary", nargs="?", const=True,
                        help="术语表：留空用内置字典，给路径则合并自定义词条")

    for name, handler, help_text in (
        ("check", command_check, "校验规格并预估功能点"),
        ("report", command_report, "输出送审自检清单"),
    ):
        node = sub.add_parser(name, parents=[common], help=help_text)
        if name == "report":
            node.add_argument("-o", "--output", help="写出到文件，缺省打到标准输出")
        node.set_defaults(handler=handler)

    build_parser = sub.add_parser("build", parents=[common], help="生成拆分表")
    build_parser.add_argument("-o", "--output", required=True, help="输出的 xlsx")
    build_parser.add_argument("--template", help="送审表模板 xlsx：沿用它的列布局与观感")
    build_parser.add_argument("--sheet", help="模板的工作表名")
    build_parser.add_argument("--l1", help="一级模块取值，默认取规格里的 project")
    build_parser.add_argument("--order", choices=("spec", "crud"), default="spec",
                              help="模块内过程顺序：spec 按声明（默认），crud 按送审增删改查惯例")
    build_parser.add_argument("--user-format", choices=sheet.USER_FORMATS, default="plain",
                              help="功能用户写法：plain 纯用户，dual 发起者/接收者两行")
    build_parser.add_argument("--require-confirmed", action="store_true",
                              help="有未标 confirmed 的过程就不写出")
    build_parser.add_argument("--force", action="store_true", help="校验不通过也写出")
    build_parser.add_argument("--selfcheck", help="同时写出自检清单 markdown")
    build_parser.set_defaults(handler=command_build)

    audit_parser = sub.add_parser("audit", help="读已有拆分表跑检查")
    audit_parser.add_argument("xlsx")
    audit_parser.add_argument("--sheet")
    audit_parser.add_argument("--max-issues", type=int, default=25)
    audit_parser.add_argument("--strict", action="store_true")
    audit_parser.add_argument("--glossary", nargs="?", const=True,
                              help="术语表：留空用内置字典，给路径则合并自定义词条")
    audit_parser.add_argument("--inventory", help=argparse.SUPPRESS)
    audit_parser.set_defaults(handler=command_audit)

    coverage_parser = sub.add_parser("coverage", help="核对接口/页面清单与规格的覆盖")
    coverage_parser.add_argument("spec")
    coverage_parser.add_argument("inventory")
    coverage_parser.add_argument("--json", action="store_true")
    coverage_parser.add_argument("-o", "--output")
    coverage_parser.set_defaults(handler=command_coverage)

    normalize_parser = sub.add_parser("normalize", help="把已有表的子过程描述改成当前句式")
    normalize_parser.add_argument("xlsx")
    normalize_parser.add_argument("--sheet")
    normalize_parser.add_argument("-o", "--output", help="改写后的新文件；不会覆盖原表")
    normalize_parser.add_argument("--write", action="store_true", help="真正写出；缺省只干跑")
    normalize_parser.add_argument("--glossary", nargs="?", const=True,
                                  help="同时按术语表清理数据组与数据属性；留空用内置字典")
    normalize_parser.add_argument("--max-issues", type=int, default=50)
    normalize_parser.set_defaults(handler=command_normalize)
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    return args.handler(args)


if __name__ == "__main__":
    sys.exit(main())
