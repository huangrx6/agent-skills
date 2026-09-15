#!/usr/bin/env python3
"""整页贴图 pptx：每页截图 → 16:9 pptx（一页一张满版图）。

**为什么是贴图而不是原生 shapes**（这是 #77 拍的板，值得写下来）：
- riso 的识别点（套色错位、颗粒、半调网点）在 pptx 原生形状里做不出原味 ——
  浏览器渲染是唯一能 100% 还原的路径
- 代价：PPT 里每页是一张图，**改不了字**；要改字得回改 `deck-spec.json` 再重出
- 要"能改字的 pptx"就得走原生 shapes 那条路，效果必然打折 —— 两条不能兼得。
  选它是因为这个 skill 的定位就是"视觉做到极近真实孔版"，不是"做个能改的普通 PPT"

跑法：python3 make_pptx.py --png-dir pages/ -o deck.pptx --width 1600 --height 900
"""
from __future__ import annotations

import argparse
import glob
import os
import sys

from pptx import Presentation
from pptx.util import Emu, Inches


def build(png_paths: list[str], out: str, width: int, height: int) -> None:
    if width <= 0 or height <= 0:
        raise SystemExit(f"✗ 版心必须是正数，收到 {width}x{height}")
    prs = Presentation()
    # 用整数运算算高度：`Emu(int(slide_width / ratio))` 先转 float 再取整，是多余的
    # 精度往返，且 height=0 时会 ZeroDivisionError（前面已拦住）。
    # `slide_width * height // width` 结果一样、全整数 —— 与 layout.py 同一手法：
    # 遇到不需要转换的地方就换个写法，而不是加一个永远不会触发的 try。
    slide_width = Inches(13.333)
    slide_height = Emu(slide_width * height // width)
    prs.slide_width = slide_width
    prs.slide_height = slide_height
    blank = prs.slide_layouts[6]                       # 空白版式，不放任何占位符
    for path in png_paths:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(path, Emu(0), Emu(0), width=slide_width, height=slide_height)
    prs.save(out)


def main(argv: list[str]) -> int:
    ap = argparse.ArgumentParser(description="截图 → pptx（一页一张满版图）")
    ap.add_argument("--png-dir", required=True)
    ap.add_argument("-o", "--out", required=True)
    ap.add_argument("--width", type=int, default=1600)
    ap.add_argument("--height", type=int, default=900)
    args = ap.parse_args(argv[1:])
    pages = sorted(glob.glob(os.path.join(args.png_dir, "page-*.png")))
    if not pages:
        raise SystemExit(f"✗ {args.png_dir} 里没有 page-*.png")
    build(pages, args.out, args.width, args.height)
    print(f"✓ 已写出 {args.out}（{len(pages)} 页 / {os.path.getsize(args.out)} 字节）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
