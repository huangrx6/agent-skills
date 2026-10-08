"""交付回归：真实图框、字体便携性、逐页 PDF 验收和 GIF 时间不能静默漂移。"""
from __future__ import annotations

import contextlib
import importlib.util
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from PIL import Image
from pptx import Presentation
from pptx.enum.text import PP_ALIGN

SCRIPTS = Path(__file__).resolve().parents[2] / "skills" / "deck-authoring" / "scripts"


def load(name):
    spec = importlib.util.spec_from_file_location("delivery_regression_" + name,
                                                 SCRIPTS / (name + ".py"))
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


fonts, native, pdf, assets, animate = (load(n) for n in
                                     ("fonts", "pptx_native", "pdf", "image_source", "animate"))
VARS = {"--paper": "#FFFFFF", "--text": "#000000", "--accent": "#0033CC",
        "--display": "Arial, Songti SC, serif", "--body": "Arial, Songti SC, serif"}


class FontEmbedding(unittest.TestCase):
    def test_file_url_and_relative_url_embed_without_nesting_style(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            font = root / "字体 with space.ttf"
            font.write_bytes(b"font-fixture")
            source = root / "source.html"
            output = root / "elsewhere" / "portable.html"
            source.write_text('<html><head><style>@font-face{font-family:"A";'
                              f'src:url("{font.as_uri()}");font-weight:700}}'
                              '@font-face{font-family:"B";src:url("字体 with space.ttf");'
                              'font-style:italic}.title{color:red}</style></head></html>')
            fonts.embed(str(source), str(output))
            font.unlink()
            text = output.read_text()
            self.assertEqual(text.count("<style>"), 1)
            self.assertEqual(text.count("</style>"), 1)
            self.assertEqual(text.count("data:font/truetype;base64,"), 2)
            self.assertNotIn("file://", text)
            self.assertIn("font-weight:700", text)
            self.assertIn("font-style:italic", text)
            self.assertIn(".title{color:red}", text)


class NativeGeometry(unittest.TestCase):
    def test_cover_and_contain_preserve_geometry_and_crop(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "portrait.png"
            Image.new("RGB", (400, 800), "red").save(path)
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            native.add_picture(slide, str(path), (900, 100, 600, 400), "cover")
            picture = slide.shapes[0]
            self.assertEqual(picture.height / native.EMU_PER_PX, 400)
            self.assertEqual(picture.width / native.EMU_PER_PX, 600)
            self.assertAlmostEqual(picture.crop_top, 1 / 3, places=4)
            self.assertAlmostEqual(picture.crop_bottom, 1 / 3, places=4)
            native.add_picture(slide, str(path), (900, 100, 600, 400), "contain")
            picture = slide.shapes[1]
            self.assertEqual(picture.width / native.EMU_PER_PX, 200)
            self.assertEqual(picture.left / native.EMU_PER_PX, 1100)
            self.assertEqual(picture.height / native.EMU_PER_PX, 400)
            self.assertEqual(picture.crop_top, 0)

    def test_object_position_controls_crop(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "portrait.png"
            Image.new("RGB", (400, 800), "red").save(path)
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            native.add_picture(slide, str(path), (0, 0, 600, 400), "cover", "50% 100%")
            self.assertAlmostEqual(slide.shapes[0].crop_top, 2 / 3, places=4)
            self.assertEqual(slide.shapes[0].crop_bottom, 0)

    def test_computed_typography_overrides_intent(self):
        measured = {"slides": [{"x": 0, "y": 0}], "elements": [{
            "id": "t", "slide": 1, "x": 50, "y": 50, "w": 800, "h": 100,
            "fontSize": 40, "fontWeight": 400, "textAlign": "center",
            "lineHeight": "48px", "letterSpacing": "-1px", "fontFamily": "Songti SC",
        }]}
        manifest = [{"id": "t", "slide": 1, "role": "title", "text": "实际字号",
                     "fontSize": 100}]
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "native.pptx"
            native._build(VARS, measured, manifest, tmp, str(out))
            p = Presentation(out).slides[0].shapes[0].text_frame.paragraphs[0]
            self.assertEqual(p.runs[0].font.size.pt, 30)
            self.assertEqual(p.alignment, PP_ALIGN.CENTER)
            self.assertEqual(p.line_spacing.pt, 36)
            self.assertEqual(p.runs[0]._r.get_or_add_rPr().get("spc"), "-75")

    def test_catalog_cjk_fonts_have_east_asian_typeface(self):
        for family in ("霞鹜文楷", "LXGW WenKai", "MiSans", "得意黑 Smiley Sans"):
            with self.subTest(family=family):
                self.assertIsNotNone(native.east_asian_family(family + ", sans-serif"))

    def test_donut_categories_remain_identifiable(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        native.add_chart(slide, {"chart": "donut", "data": [
            {"label": "A", "value": 70}, {"label": "B", "value": 30}]},
                         (0, 0, 600, 400), VARS)
        chart = slide.shapes[0].chart
        self.assertTrue(chart.has_legend)
        points = chart.series[0].points
        self.assertNotEqual(points[0].format.fill.fore_color.rgb,
                            points[1].format.fill.fore_color.rgb)

    def test_combo_is_not_silently_changed_to_columns(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        with self.assertRaisesRegex(SystemExit, "combo"):
            native.add_chart(slide, {"chart": "combo", "series": [
                {"name": "A", "data": [{"label": "Q1", "value": 1}]}]},
                             (0, 0, 600, 400), VARS)


class NativeChartData(unittest.TestCase):
    def chart(self, spec):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        native.add_chart(slide, spec, (0, 0, 600, 400), VARS)
        return slide.shapes[0].chart

    def test_scatter_accepts_single_series_without_losing_points(self):
        chart = self.chart({"chart": "scatter", "series": [{"name": "S", "data": [
            {"label": "A", "x": 1, "y": 2}, {"label": "B", "x": 3, "y": 4}]}]})
        self.assertEqual([value for series in chart.series for value in series.values], [2, 4])
        self.assertEqual([float(node.text) for node in chart._chartSpace.xpath(
            ".//c:xVal/c:numRef/c:numCache/c:pt/c:v")], [1, 3])

    def test_scatter_explicit_y_has_priority_over_legacy_value(self):
        chart = self.chart({"chart": "scatter", "data": [
            {"label": "A", "x": 1, "y": 2, "value": 999}]})
        self.assertEqual(chart.series[0].values, (2.0,))

    def test_series_align_by_category_and_missing_values_stay_gaps(self):
        chart = self.chart({"chart": "line", "series": [
            {"name": "S1", "data": [{"label": "A", "value": 1}, {"label": "B", "value": 2}]},
            {"name": "S2", "data": [{"label": "B", "value": 20}, {"label": "A", "value": 10},
                                     {"label": "C", "value": 0}]},
            {"name": "S3", "data": [{"label": "D", "value": 7}]},
        ]})
        self.assertEqual([category.label for category in chart.plots[0].categories], ["A", "B", "C", "D"])
        self.assertEqual(chart.series[0].values, (1.0, 2.0, None, None))
        self.assertEqual(chart.series[1].values, (10.0, 20.0, 0.0, None))
        self.assertEqual(chart.series[2].values, (None, None, None, 7.0))
        self.assertEqual(chart._chartSpace.xpath(".//c:dispBlanksAs")[0].get("val"), "gap")

    def test_duplicate_category_cannot_overwrite_an_observation(self):
        with self.assertRaisesRegex(SystemExit, "重复"):
            self.chart({"chart": "line", "series": [{"name": "S", "data": [
                {"label": "A", "value": 1}, {"label": "A", "value": 2}]}]})

    def test_single_series_donut_keeps_all_categories_and_valid_colors(self):
        chart = self.chart({"chart": "donut", "series": [{"name": "S", "data": [
            {"label": "A", "value": 50}, {"label": "B", "value": 30},
            {"label": "C", "value": 20}]}]})
        self.assertEqual(chart.series[0].values, (50.0, 30.0, 20.0))
        self.assertEqual(len({str(point.format.fill.fore_color.rgb)
                              for point in chart.series[0].points}), 3)

    def test_unit_labels_preserve_fractional_and_tiny_values(self):
        chart = self.chart({"chart": "bar", "unit": "ms", "data": [
            {"label": "fraction", "value": 0.125}, {"label": "tiny", "value": 1e-12}]})
        self.assertEqual(chart.series[0].values, (0.125, 1e-12))
        self.assertEqual(chart.plots[0].data_labels.number_format, 'General"ms"')
        self.assertFalse(chart.plots[0].data_labels.number_format_is_linked)

    def test_line_markers_and_area_outlines_write_explicit_palette_colors(self):
        expected = ["0033CC", "9EB1EC"]
        series = [
            {"name": "S1", "data": [{"label": "A", "value": 1}, {"label": "B", "value": 2}]},
            {"name": "S2", "data": [{"label": "A", "value": 3}, {"label": "B", "value": 4}]},
        ]
        for kind in ("line", "area"):
            with self.subTest(kind=kind):
                chart = self.chart({"chart": kind, "series": series})
                prefix = f".//c:{kind}Chart/c:ser"
                for tail in ("/c:spPr/a:solidFill/a:srgbClr",
                             "/c:spPr/a:ln/a:solidFill/a:srgbClr"):
                    self.assertEqual([node.get("val") for node in chart._chartSpace.xpath(prefix + tail)],
                                     expected)
                if kind == "line":
                    for tail in ("/c:marker/c:spPr/a:solidFill/a:srgbClr",
                                 "/c:marker/c:spPr/a:ln/a:solidFill/a:srgbClr"):
                        self.assertEqual([node.get("val") for node in chart._chartSpace.xpath(prefix + tail)],
                                         expected)


class PdfPageValidation(unittest.TestCase):
    def test_wrong_height_mixed_sizes_and_missing_boxes_are_rejected(self):
        for boxes, pages in ((["0 0 1200 100"], 1),
                             (["0 0 1200 675", "0 0 612 792"], 2), ([], 1)):
            with self.subTest(boxes=boxes), contextlib.redirect_stdout(io.StringIO()):
                code = pdf.report({"bytes": 2000, "pages": pages, "expected_pages": pages,
                                   "images": 0, "fonts": 1, "mediaboxes": boxes}, "test.pdf")
                self.assertEqual(code, 1)


class ImageResolution(unittest.TestCase):
    def check_image(self, width, height, slots):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            spec = root / "deck.spec.json"
            spec.write_text(json.dumps({"deck": {"slides": [
                {"image": "picture.png", "visual": {"ratio": "1:1"}}]}}))
            Image.new("RGB", (width, height), "red").save(root / "picture.png")
            geometry = {"slides": {i + 1: {"file": "picture.png", "box": box}
                                   for i, box in enumerate(slots)}}
            with patch.object(assets, "_slot_geometry", return_value=(geometry, "test")):
                return assets.check_images(str(spec), tmp)

    def test_cover_checks_height_after_crop(self):
        code, problems, _ = self.check_image(1280, 64, [{"w": 600, "h": 400, "objectFit": "cover"}])
        self.assertEqual(code, 1)
        self.assertIn("64px 高", problems[0])

    def test_contain_checks_only_the_displayed_image_size(self):
        code, problems, _ = self.check_image(1280, 64, [{"w": 600, "h": 400, "objectFit": "contain"}])
        self.assertEqual(code, 0, problems)

    def test_reused_image_must_satisfy_largest_slot(self):
        code, problems, _ = self.check_image(1280, 900, [{"w": 400, "h": 300},
                                                      {"w": 1400, "h": 800}])
        self.assertEqual(code, 1)
        self.assertIn("第 2 页", problems[0])


class GifTiming(unittest.TestCase):
    def test_ten_seconds_at_24fps_does_not_shrink_to_9_6_seconds(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            frames = []
            for i in range(240):
                path = root / f"{i}.png"
                im = Image.new("RGB", (16, 16), (i, i * 7 % 256, i * 17 % 256))
                im.putpixel((i % 16, i // 16), (255, 255, 255))
                im.save(path)
                frames.append(str(path))
            out = root / "animation.gif"
            animate.encode_gif(frames, str(out), 24)
            info = animate.inspect_media(str(out), 24)
            self.assertLessEqual(abs(info["gif_ms"] - 10000), 10)


if __name__ == "__main__":
    unittest.main()
