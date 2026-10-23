"""模板探测、缺列补齐与回读。

跑法：
    python3 -m unittest discover -s tests/cosmic-split -v

布局不假设是这个 skill 最核心的一条：上一版把列号写死在文档里，第一次拿到真实模板就错位。
所以这里把实测见过的三种模板形态都构造成夹具，钉住探测行为。
需要 openpyxl；缺失时整个模块跳过。
"""

from __future__ import annotations

import os
import sys
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.join(os.path.dirname(os.path.dirname(HERE)), "skills", os.path.basename(HERE))
SCRIPTS = os.path.join(SKILL, "scripts")
sys.path.insert(0, SCRIPTS)

try:
    import openpyxl

    HAVE_OPENPYXL = True
except ImportError:  # pragma: no cover - 环境缺依赖
    HAVE_OPENPYXL = False

import builder  # noqa: E402
import sheet  # noqa: E402

from test_pipeline import FIXTURE  # noqa: E402


def _rows() -> list[builder.Row]:
    return builder.build(__import__("spec").load_from_dict(FIXTURE)).rows


@unittest.skipUnless(HAVE_OPENPYXL, "需要 openpyxl")
class SheetTestCase(unittest.TestCase):
    def setUp(self):
        self.tmp = os.path.join(
            os.path.dirname(os.path.abspath(__file__)), "_tmp_sheet"
        )
        os.makedirs(self.tmp, exist_ok=True)
        self.rows = _rows()

    def tearDown(self):
        for name in os.listdir(self.tmp):
            os.remove(os.path.join(self.tmp, name))
        os.rmdir(self.tmp)

    def _path(self, name: str) -> str:
        return os.path.join(self.tmp, name)

    def _make(self, name: str, headers: list[str], data: list[list], header_row: int = 1,
              merges: list[str] | None = None, sheet_title: str = "Sheet1") -> str:
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = sheet_title
        for offset, label in enumerate(headers):
            if label is not None:
                worksheet.cell(row=header_row, column=offset + 1, value=label)
        for row_offset, row in enumerate(data):
            for column, value in enumerate(row, start=1):
                if value is not None:
                    worksheet.cell(row=header_row + 1 + row_offset, column=column, value=value)
        for rng in merges or []:
            worksheet.merge_cells(rng)
        path = self._path(name)
        workbook.save(path)
        return path


class TestInspect(SheetTestCase):
    def test_标准十二列_表头在第一行(self):
        path = self._make("std.xlsx", [
            "一级模块", "二级模块", "三级模块", "功能用户", "触发事件", "功能过程",
            "子过程描述", "数据移动类型", "数据组", "数据属性", "复用度", "CFP",
        ], [["项目", "模块", "子模块", "用户", "事件", "功能", "描述", "E", "组", "属性", 1, 2]])
        template = sheet.inspect(path)
        self.assertEqual(template.header_row, 1)
        self.assertEqual(template.layout["l1"], 1)
        self.assertEqual(template.layout["cfp"], 12)
        self.assertEqual(template.layout["move"], 8)

    def test_四列极简表_移动类型列叫标识(self):
        # 真实表：二级模块 / 三级模块 / 功能过程 / 标识，没有 CFP 列也没有合并
        path = self._make("mini.xlsx", ["二级模块", "三级模块", "功能过程", "标识"], [
            ["用户认证", "账号登录", "管理员登录", "E"],
            ["用户认证", "账号登录", "管理员登录", "W"],
            ["用户认证", "账号登录", "管理员登录", "X"],
        ])
        template = sheet.inspect(path)
        self.assertEqual(template.layout["l2"], 1)
        self.assertEqual(template.layout["l3"], 2)
        self.assertEqual(template.layout["func"], 3)
        self.assertEqual(template.layout.get("move"), 4, "「标识」应被识别为移动类型列")
        self.assertNotIn("cfp", template.layout)
        self.assertTrue(any("CFP" in note for note in template.notes))

    def test_通用标识列装别的值时不采纳(self):
        # 「标识」太通用，若那一列装的不是 E/R/W/X 就不该被当成移动类型列。
        path = self._make("decoy.xlsx", ["二级模块", "三级模块", "功能过程", "标识", "数据移动类型"], [
            ["模块", "子模块", "功能", "启用", "E"],
            ["模块", "子模块", "功能", "停用", "W"],
        ])
        template = sheet.inspect(path)
        self.assertEqual(template.layout.get("move"), 5, "应退到真正装着 E/R/W/X 的那一列")

    def test_多sheet表_表头在第四行(self):
        workbook = openpyxl.Workbook()
        cover = workbook.active
        cover.title = "使用说明"
        cover.cell(row=1, column=1, value="填写与核验说明")
        data = workbook.create_sheet("COSMIC拆分表")
        data.cell(row=1, column=1, value="某某平台 COSMIC 拆分表")
        data.cell(row=2, column=1, value="每个功能过程以 E 开始，以 W 或 X 结束")
        for column, label in enumerate(
            ["序号", "一级业务域", "二级功能模块", "功能用户", "触发事件", "功能过程",
             "子过程序号", "子过程", "数据移动类型", "数据组", "数据属性", "代码依据"], start=1
        ):
            data.cell(row=4, column=column, value=label)
        data.cell(row=5, column=2, value="多租户与账号")
        data.cell(row=5, column=9, value="E")
        path = self._path("multi.xlsx")
        workbook.save(path)

        template = sheet.inspect(path)
        self.assertEqual(template.sheet_name, "COSMIC拆分表", "应跳过只有说明的封面表")
        self.assertEqual(template.header_row, 4)
        self.assertEqual(template.layout["l1"], 2)
        self.assertEqual(template.layout["l2"], 3)
        self.assertEqual(template.layout["desc"], 8, "「子过程」应映射到子过程描述")
        self.assertEqual(template.layout["move"], 9)
        self.assertNotIn("l3", template.layout)
        self.assertTrue(any("没认出来" in note for note in template.notes))
        self.assertEqual(template.freeze, "A5", "冻结位置跟着表头行走")

    def test_指定sheet(self):
        path = self._make("one.xlsx", ["二级模块", "三级模块", "功能过程", "标识"],
                          [["模块", "子模块", "功能", "E"]], sheet_title="自定义表名")
        template = sheet.inspect(path, sheet="自定义表名")
        self.assertEqual(template.sheet_name, "自定义表名")

    def test_文件不存在时报清楚(self):
        with self.assertRaises(SystemExit) as ctx:
            sheet.inspect(self._path("nope.xlsx"))
        self.assertIn("读不到模板文件", str(ctx.exception))


class TestReconcile(SheetTestCase):
    def test_缺列追加到末尾而不是丢数据(self):
        path = self._make("mini.xlsx", ["二级模块", "三级模块", "功能过程", "标识"],
                          [["模块", "子模块", "功能", "E"]])
        template = sheet.reconcile(sheet.inspect(path))
        for logical in ("l1", "user", "event", "desc", "group", "attrs", "cfp"):
            self.assertIn(logical, template.layout, f"{logical} 应被追加")
        self.assertEqual(template.appended["l1"], 5, "追加列应接在最右侧")
        self.assertTrue(
            any("追加到末尾" in note for note in template.notes),
            "追加了哪些列必须写进备注",
        )

    def test_复用度是常量列_模板没有就不追加(self):
        path = self._make("mini.xlsx", ["二级模块", "三级模块", "功能过程", "标识"],
                          [["模块", "子模块", "功能", "E"]])
        template = sheet.reconcile(sheet.inspect(path))
        self.assertNotIn("reuse", template.appended)
        self.assertTrue(any("未追加" in note for note in template.notes))

    def test_十二列模板不需要追加(self):
        template = sheet.reconcile(sheet.Template(l1_value="项目"))
        self.assertEqual(template.appended, {})


class TestWriteAndReadBack(SheetTestCase):
    def test_写出后能原样读回(self):
        template = sheet.Template(l1_value="回归项目")
        path = sheet.write(self.rows, self._path("out.xlsx"), template)

        read_back = __import__("cosmic").read_existing_table(path)
        self.assertEqual(len(read_back), len(self.rows))
        for original, restored in zip(self.rows, read_back):
            self.assertEqual(restored.func, original.func)
            self.assertEqual(restored.desc, original.desc)
            self.assertEqual(restored.move, original.move)
            self.assertEqual(restored.attrs, original.attrs)
            self.assertEqual(restored.group, original.group)

    def test_回读时CFP按块行数回填(self):
        # 不能在读到 E 行时就算 CFP —— 那一刻这个块的序列还没读完。
        template = sheet.Template(l1_value="回归项目")
        path = sheet.write(self.rows, self._path("out.xlsx"), template)
        read_back = __import__("cosmic").read_existing_table(path)
        heads = [row for row in read_back if row.is_process_head]
        self.assertEqual(len(heads), 2)
        for head in heads:
            self.assertGreaterEqual(head.fp, 2)
        self.assertEqual(sum(head.fp for head in heads), len(read_back))

    def test_回读的表能过校验(self):
        rows = __import__("cosmic").read_existing_table(
            sheet.write(self.rows, self._path("out.xlsx"), sheet.Template(l1_value="回归项目"))
        )
        issues = __import__("checks").check(rows)
        self.assertEqual([issue for issue in issues if "CFP" not in issue], [])

    def test_功能用户列按连续相同取值合并(self):
        path = sheet.write(self.rows, self._path("out.xlsx"),
                           sheet.Template(l1_value="回归项目"))
        worksheet = openpyxl.load_workbook(path).active
        user_ranges = [
            rng for rng in worksheet.merged_cells.ranges
            if rng.min_col == sheet.DEFAULT_LAYOUT["user"]
        ]
        # 运营看板里报表查询是业务用户、采集入库是定时任务，必然分成两块以上
        self.assertGreaterEqual(len(user_ranges), 2)

    def test_缺一级模块取值时报错(self):
        with self.assertRaises(SystemExit) as ctx:
            sheet.write(self.rows, self._path("out.xlsx"), sheet.Template())
        self.assertIn("一级模块取值", str(ctx.exception))


if __name__ == "__main__":
    unittest.main()
