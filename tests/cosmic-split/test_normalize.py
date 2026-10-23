"""已有表的安全改写：只改文字，保留合并、样式与其他表。"""

from __future__ import annotations

import contextlib
import io
import json
import os
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

HERE = Path(__file__).resolve().parent
SCRIPTS = HERE.parents[1] / 'skills/cosmic-split/scripts'
sys.path.insert(0, str(SCRIPTS))

import cosmic  # noqa: E402
import openpyxl  # noqa: E402
import sheet  # noqa: E402

# 静态分析器看不到 sys.path 挂载的同级模块。
# pyright: reportMissingImports=false


def active(workbook):
    """openpyxl 的类型里 workbook.active 可能为空；测试里直接取默认表。"""
    worksheet = workbook.active
    if worksheet is None:  # pragma: no cover - 由 openpyxl 保证
        raise AssertionError("工作簿没有默认工作表")
    return worksheet


def run(*argv: str) -> tuple[int, str, str]:
    out, err = io.StringIO(), io.StringIO()
    with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
        try:
            code = cosmic.main(list(argv))
        except SystemExit as exc:
            code = exc.code if isinstance(exc.code, int) else 1
    return code, out.getvalue(), err.getvalue()


LEGACY = {
    2: "接收地址新增请求并校验输入数据",
    3: "从数据存储读取地址数据",
    4: "将地址数据写入数据存储",
    5: "向用户返回地址新增结果",
}


class NormalizeTests(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp(prefix="cosmic-normalize-")
        self.addCleanup(shutil.rmtree, self.dir, True)
        self.table = os.path.join(self.dir, "legacy.xlsx")
        spec_path = os.path.join(self.dir, "spec.json")
        self.assertEqual(run("init", spec_path)[0], 0)
        self.assertEqual(run("build", spec_path, "-o", self.table)[0], 0)
        self.edit_legacy()

    def edit_legacy(self):
        workbook = openpyxl.load_workbook(self.table)
        worksheet = active(workbook)
        layout = sheet.inspect(self.table).layout
        for row, text in LEGACY.items():
            worksheet.cell(row=row, column=layout['desc'], value=text)
        workbook.save(self.table)
        workbook.close()

    def descs(self, path):
        workbook = openpyxl.load_workbook(path)
        worksheet = active(workbook)
        column = sheet.inspect(path).layout['desc']
        values = [worksheet.cell(row=row, column=column).value for row in range(2, 6)]
        workbook.close()
        return values

    def test_干跑列出改动但不写文件(self):
        output = os.path.join(self.dir, "new.xlsx")
        code, out, _err = run("normalize", self.table)
        self.assertEqual(code, 0)
        self.assertIn("干跑", out)
        self.assertIn("输入地址新增请求", out)
        self.assertFalse(os.path.exists(output))
        self.assertEqual(self.descs(self.table)[0], LEGACY[2], "干跑不能改原表")

    def test_写出后原表不变新表改写(self):
        output = os.path.join(self.dir, "new.xlsx")
        code, out, _err = run("normalize", self.table, "-o", output, "--write")
        self.assertEqual(code, 0, out)
        self.assertIn("4 处改写", out)
        self.assertEqual(self.descs(self.table)[0], LEGACY[2])
        self.assertEqual(self.descs(output), [
            "输入地址新增请求", "读取地址数据", "保存地址数据", "返回地址新增结果",
        ])
        self.assertEqual(run("audit", output)[0], 0)

    def test_保留其他工作表与合并单元格(self):
        workbook = openpyxl.load_workbook(self.table)
        workbook.create_sheet("自检清单")["A1"] = "原样保留"
        workbook.save(self.table)
        opened = openpyxl.load_workbook(self.table)
        before = len(active(opened).merged_cells.ranges)
        opened.close()

        output = os.path.join(self.dir, "new.xlsx")
        self.assertEqual(run("normalize", self.table, "-o", output, "--write")[0], 0)
        workbook = openpyxl.load_workbook(output)
        self.assertIn("自检清单", workbook.sheetnames)
        self.assertEqual(workbook["自检清单"]["A1"].value, "原样保留")
        self.assertEqual(len(active(workbook).merged_cells.ranges), before)
        workbook.close()

    def test_术语表清理英文但不拆字段标识(self):
        workbook = openpyxl.load_workbook(self.table)
        worksheet = active(workbook)
        layout = sheet.inspect(self.table).layout
        worksheet.cell(row=2, column=layout['desc'], value="Download状态查询")
        worksheet.cell(row=2, column=layout['attrs'], value="用户ID,addressID")
        workbook.save(self.table)
        workbook.close()

        output = os.path.join(self.dir, "glossary.xlsx")
        code, out, _err = run("normalize", self.table, "-o", output, "--write", "--glossary")
        self.assertEqual(code, 0, out)
        workbook = openpyxl.load_workbook(output)
        worksheet = active(workbook)
        self.assertEqual(worksheet.cell(row=2, column=layout['desc']).value, "下载状态查询")
        self.assertEqual(worksheet.cell(row=2, column=layout['attrs']).value, "用户编码,addressID")
        workbook.close()

    def test_不覆盖原表且_write必须给输出(self):
        code, _out, err = run("normalize", self.table, "-o", self.table, "--write")
        self.assertEqual(code, 1)
        self.assertIn("路径冲突", err)
        code, _out, err = run("normalize", self.table, "--write")
        self.assertEqual(code, 1)
        self.assertIn("需要同时指定 -o", err)

    def test_无需改写时明确说明(self):
        output = os.path.join(self.dir, "clean.xlsx")
        self.assertEqual(run("normalize", self.table, "-o", output, "--write")[0], 0)
        code, out, _err = run("normalize", output)
        self.assertEqual(code, 0)
        self.assertIn("没有需要改写", out)

    def test_指定工作表写回同一张表(self):
        workbook = openpyxl.load_workbook(self.table)
        first = active(workbook)
        first.title = '保留表'
        selected = workbook.copy_worksheet(first)
        selected.title = '目标表'
        selected['G2'] = '接收退款申请请求并校验输入数据'
        workbook.save(self.table); workbook.close()
        output = os.path.join(self.dir, 'selected.xlsx')
        self.assertEqual(run('normalize', self.table, '--sheet', '目标表', '-o', output, '--write')[0], 0)
        workbook = openpyxl.load_workbook(output)
        self.assertEqual(workbook['保留表']['G2'].value, LEGACY[2])
        self.assertEqual(workbook['目标表']['G2'].value, '输入退款申请请求')
        workbook.close()

    def test_公式和未识别的业务文字原样保留(self):
        workbook = openpyxl.load_workbook(self.table)
        worksheet = active(workbook)
        worksheet['G2'] = '=IF(TRUE,"Download","ID")'
        worksheet['G3'] = '  读取查询查询记录\n并保留原文  '
        workbook.save(self.table); workbook.close()
        output = os.path.join(self.dir, 'formulas.xlsx')
        self.assertEqual(run('normalize', self.table, '-o', output, '--write', '--glossary')[0], 0)
        workbook = openpyxl.load_workbook(output)
        self.assertEqual(active(workbook)['G2'].value, '=IF(TRUE,"Download","ID")')
        self.assertEqual(active(workbook)['G2'].data_type, 'f')
        self.assertEqual(active(workbook)['G3'].value, '  读取查询查询记录\n并保留原文  ')
        workbook.close()

    def test_术语改写超限不覆盖已有输出(self):
        glossary = Path(self.dir) / 'glossary.json'
        glossary.write_text(json.dumps({'map': {'地址': '字' * 32768}}))
        output = Path(self.dir) / 'existing.xlsx'
        output.write_bytes(b'existing')
        self.assertEqual(run('normalize', self.table, '-o', str(output), '--write',
                             '--glossary', str(glossary))[0], 1)
        self.assertEqual(output.read_bytes(), b'existing')


if __name__ == "__main__":
    unittest.main()
