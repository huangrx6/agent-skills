"""独立业务反例与审计回归，验证原始字段、来源和过程边界不被生成器改写。"""
import copy
import json
import tempfile
import unittest
from pathlib import Path

from test_cli import run
import builder
import checks
import cosmic
import sheet
import spec
import openpyxl


def fixture():
    return copy.deepcopy(cosmic.STARTER_SPEC)


def build(value):
    return builder.build(spec.load_from_dict(value))


class MeasurementTests(unittest.TestCase):
    def test_新增可包含读取和返回(self):
        result = build(fixture())
        self.assertEqual([r.move for r in result.rows], ['E', 'R', 'W', 'X'])
        self.assertEqual(result.total_cfp, 4)
        self.assertEqual(checks.check(result.rows), [])

    def test_同一过程多个输入不增加过程数(self):
        value = fixture()
        value['processes'][0]['movements'].insert(1, {
            'type': 'E', 'object': '联系人', 'group': '联系人资料',
            'attributes': ['姓名'], 'description': '接收联系人姓名'})
        result = build(value)
        self.assertEqual(result.process_count, 1)
        self.assertEqual(result.total_cfp, 5)
        self.assertEqual(len(checks.process_blocks(result.rows)), 1)
        self.assertEqual(checks.check(result.rows), [])

    def test_单字段及英文业务标识合法(self):
        value = fixture()
        value['processes'][0]['movements'][0]['attributes'] = ['addressID']
        result = build(value)
        self.assertEqual(result.rows[0].attrs, 'addressID')
        self.assertEqual(checks.check(result.rows), [])

    def test_相同数据组可跨过程复用(self):
        value = fixture()
        second = copy.deepcopy(value['processes'][0])
        second.update(id='FP-002', name='地址修改', trigger='用户修改地址')
        value['processes'].append(second)
        result = build(value)
        self.assertEqual(checks.check(result.rows), [])
        self.assertEqual(result.rows[0].attrs, result.rows[4].attrs)

    def test_重复移动需要需求理由(self):
        value = fixture()
        proc = value['processes'][0]
        proc['movements'].insert(2, copy.deepcopy(proc['movements'][1]))
        self.assertTrue(any('重复移动需需求依据' in x for x in checks.check(build(value).rows)))
        proc['movements'][2]['repeat_reason'] = '需求 R2 明确要求用户确认后再次读取最新地址'
        self.assertEqual(checks.check(build(value).rows), [])

    def test_已确认必须有来源(self):
        value = fixture()
        value['processes'][0]['status'] = 'confirmed'
        with self.assertRaises(spec.SpecError):
            build(value)
        value['processes'][0]['source'] = '需求说明.md#地址管理'
        info = checks.summary(build(value).rows)
        self.assertEqual((info['confirmed_cfp'], info['unverified_cfp']), (4, 0))

    def test_逐移动来源可替代过程来源(self):
        value = fixture()
        value['processes'][0]['status'] = 'confirmed'
        for movement in value['processes'][0]['movements']:
            movement['source'] = '需求说明.md#新增地址'
        self.assertEqual(checks.summary(build(value).rows)['confirmed_cfp'], 4)

    def test_来源标记不同分别汇总(self):
        value = fixture()
        second = copy.deepcopy(value['processes'][0])
        second.update(id='FP-002', name='地址修改', status='confirmed', source='需求 R2')
        value['processes'].append(second)
        info = checks.summary(build(value).rows)
        self.assertEqual((info['confirmed_cfp'], info['unverified_cfp']), (4, 4))

    def test_改变目标不改变计量(self):
        value = fixture()
        before = build(value).rows
        value['target_cfp'] = [700, 800]
        self.assertEqual(build(value).rows, before)

    def test_不保留旧规格或检查模式(self):
        for field, value in [('version', 1), ('entities', {}), ('profile', 'standard')]:
            with self.subTest(field=field):
                data = fixture()
                data[field] = value
                with self.assertRaises(spec.SpecError):
                    build(data)

    def test_字段全量保留不补造(self):
        value = fixture()
        fields = [f'业务字段{n}' for n in range(12)]
        value['processes'][0]['movements'][0]['attributes'] = fields
        self.assertEqual(build(value).rows[0].attrs.split(','), fields)

    def test_不支持的版本与错误字段类型被拒绝(self):
        for mutation in (
            lambda s: s.update(version=2),
            lambda s: s['processes'][0]['movements'][0].update(attributes=[None]),
            lambda s: s['processes'][0]['movements'][0].update(attributes=['名称', '名称']),
            lambda s: s['processes'][0].update(user='范围外用户'),
            lambda s: s['processes'][0].update(status='done'),
        ):
            with self.subTest(mutation=mutation):
                value = fixture(); mutation(value)
                with self.assertRaises(spec.SpecError):
                    build(value)

    def test_审批按显式移动计数(self):
        value = fixture()
        p = value['processes'][0]
        p.update(name='申请审批')
        p['movements'] = p['movements'][:3]
        self.assertEqual([r.move for r in build(value).rows], ['E', 'R', 'W'])
        self.assertEqual(build(value).total_cfp, 3)

    def test_标识首尾空白不能绕过重复判断(self):
        for mutate in (
            lambda s: s['processes'][0]['movements'][1].update(object='地址 '),
            lambda s: s['processes'][0].update(id=' FP-001'),
            lambda s: s['processes'][0].update(name='地址新增 '),
            lambda s: s['modules'][0].update(l2='模块 '),
        ):
            with self.subTest(mutate=mutate):
                value = fixture()
                mutate(value)
                with self.assertRaisesRegex(spec.SpecError, '首尾空白'):
                    build(value)

    def test_加载不修改原始对象(self):
        value = fixture(); original = copy.deepcopy(value)
        build(value)
        self.assertEqual(value, original)


class AuditRegressions(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix='cosmic-regression-')
        self.root = Path(self.tmp.name)
        self.path = self.root / 'out.xlsx'
        self.rows = build(fixture()).rows
        sheet.write(self.rows, self.path, sheet.Template(l1_value='项目'), trace=True)

    def tearDown(self):
        self.tmp.cleanup()

    def edit(self, cell, value):
        wb = openpyxl.load_workbook(self.path)
        wb.active[cell] = value
        wb.save(self.path); wb.close()

    def issues(self):
        return checks.check(cosmic.read_existing_table(self.path))

    def test_原表999不能通过(self):
        self.edit('L2', 999)
        self.assertTrue(any('声明 999' in x for x in self.issues()))
        self.assertEqual(run('audit', str(self.path))[0], 1)

    def test_非法移动不能静默删除(self):
        self.edit('H3', 'Q')
        rows = cosmic.read_existing_table(self.path)
        self.assertEqual(len(rows), 4)
        self.assertTrue(any('非法数据移动类型 第 3 行' in x for x in checks.check(rows)))

    def test_空移动也报错(self):
        self.edit('H3', None)
        self.assertTrue(any('非法数据移动类型' in x for x in self.issues()))

    def test_全列非法仍能识别明确表头(self):
        wb = openpyxl.load_workbook(self.path)
        for r in range(2, 6): wb.active.cell(r, 8, 'Q')
        wb.save(self.path); wb.close()
        self.assertEqual(sheet.inspect(self.path).layout['move'], 8)
        self.assertEqual(run('audit', str(self.path))[0], 1)

    def test_自动选表与读取一致(self):
        wb = openpyxl.load_workbook(self.path)
        wb.create_sheet('使用说明', 0)['A1'] = '说明'
        wb.save(self.path); wb.close()
        self.assertEqual(len(cosmic.read_existing_table(self.path)), 4)
        self.assertEqual(run('audit', str(self.path))[0], 0)

    def test_空模板不重复追加移动列(self):
        wb = openpyxl.Workbook()
        for key, col in sheet.DEFAULT_LAYOUT.items():
            wb.active.cell(1, col, sheet.HEADER_LABELS[key])
        path = self.root / 'empty-template.xlsx'; wb.save(path); wb.close()
        t = sheet.reconcile(sheet.inspect(path))
        self.assertEqual(t.layout['move'], 8)
        self.assertEqual(t.appended, {})

    def test_缺失CFP与公式CFP不能当作核验通过(self):
        for value in (None, '=2+2', '四', -1, 2.5):
            with self.subTest(value=value):
                self.edit('L2', value)
                self.assertTrue(any('CFP' in x for x in self.issues()))

    def test_逐行CFP求和可核对(self):
        wb = openpyxl.load_workbook(self.path)
        wb.active.unmerge_cells('L2:L5')
        for r in range(2, 6): wb.active.cell(r, 12, 1)
        wb.save(self.path); wb.close()
        self.assertEqual(self.issues(), [])

    def test_多个E写出回读仍为一个过程(self):
        value = fixture()
        value['processes'][0]['movements'].insert(1, {
            'type': 'E', 'object': '联系人', 'group': '联系人资料', 'attributes': ['姓名'], 'description': '接收联系人'})
        rows = build(value).rows
        sheet.write(rows, self.path, sheet.Template(l1_value='项目'))
        restored = cosmic.read_existing_table(self.path)
        self.assertEqual(checks.summary(restored)['processes'], 1)
        self.assertEqual(checks.summary(restored)['cfp'], 5)
        self.assertEqual(checks.check(restored), [])

    def test_追溯字段回读保留(self):
        value = fixture(); proc = value['processes'][0]
        proc.update(status='confirmed', source='需求.md#新增地址')
        rows = build(value).rows
        sheet.write(rows, self.path, sheet.Template(l1_value='项目'), trace=True)
        restored = cosmic.read_existing_table(self.path)
        for a, b in zip(rows, restored):
            self.assertEqual((a.process_id, a.status, a.source, a.object), (b.process_id, b.status, b.source, b.object))
        self.assertEqual(checks.summary(restored)['confirmed_cfp'], 4)

    def test_过程内混合核实状态不能通过(self):
        t = sheet.inspect(self.path)
        wb = openpyxl.load_workbook(self.path)
        wb.active.cell(3, t.layout['status'], 'confirmed')
        wb.active.cell(3, t.layout['source'], '需求 R1')
        wb.save(self.path); wb.close()
        self.assertTrue(any('核实状态不一致' in x for x in self.issues()))

    def test_从第四行表头读取且忽略上方说明(self):
        sheet.write(self.rows, self.path, sheet.Template(l1_value='项目', header_row=4))
        wb = openpyxl.load_workbook(self.path)
        wb.active['F2'] = '上方样例'; wb.active['H2'] = 'E'
        wb.save(self.path); wb.close()
        self.assertEqual(len(cosmic.read_existing_table(self.path)), 4)
        self.assertEqual(self.issues(), [])

    def test_init生成唯一规格并自动生成来源报告(self):
        path = self.root / 'spec.json'
        self.assertEqual(run('init', str(path))[0], 0)
        self.assertEqual(json.loads(path.read_text())['version'], 3)
        self.assertEqual(run('build', str(path), '-o', str(self.path))[0], 0)
        report = Path(str(self.path) + '.selfcheck.md').read_text()
        self.assertIn('待核实 4 CFP', report)
        self.assertIn('FP-001', report)
        self.assertNotIn('2026 年送审被打回', report)

    def test_没有独立触发的多表写入不被拆过程(self):
        value = fixture()
        value['processes'][0]['movements'].insert(3, {
            'type': 'W', 'object': '地址标签关联', 'group': '地址标签',
            'attributes': ['地址编码', '标签编码'], 'description': '保存标签关联'})
        result = build(value)
        self.assertEqual(result.process_count, 1)
        self.assertEqual(result.total_cfp, 5)
        self.assertEqual(checks.check(result.rows), [])
