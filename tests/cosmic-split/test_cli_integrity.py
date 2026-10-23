"""写出路径冲突必须在任何文件改变之前被拒绝。"""
import json
import os
import tempfile
import unittest
from pathlib import Path

from test_cli import run
import cosmic
import openpyxl


class OutputPathTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix='cosmic-paths-')
        self.root = Path(self.tmp.name)
        self.spec = self.root / 'spec.json'
        self.spec.write_text(json.dumps(cosmic.STARTER_SPEC, ensure_ascii=False), encoding='utf-8')
        self.original = self.spec.read_bytes()
        self.output = self.root / 'out.xlsx'

    def tearDown(self):
        self.tmp.cleanup()

    def test_输出与来源报告同路径时不写任何文件(self):
        self.output.write_bytes(b'existing output')
        code, _, err = run('build', str(self.spec), '-o', str(self.output), '--selfcheck', str(self.output))
        self.assertEqual(code, 1)
        self.assertIn('路径冲突', err)
        self.assertEqual(self.output.read_bytes(), b'existing output')
        self.assertEqual(self.spec.read_bytes(), self.original)

    def test_表格或报告不能覆盖规格(self):
        for args in (
            ['build', str(self.spec), '-o', str(self.spec)],
            ['build', str(self.spec), '-o', str(self.output), '--selfcheck', str(self.spec)],
            ['report', str(self.spec), '-o', str(self.spec)],
        ):
            with self.subTest(args=args):
                self.assertEqual(run(*args)[0], 1)
                self.assertEqual(self.spec.read_bytes(), self.original)
                self.assertFalse(self.output.exists())

    def test_表格和报告不能覆盖模板(self):
        template = self.root / 'template.xlsx'
        wb = openpyxl.Workbook(); wb.active.append(['功能过程', '数据移动类型']); wb.save(template); wb.close()
        original = template.read_bytes()
        for args in (
            ['build', str(self.spec), '-o', str(template), '--template', str(template)],
            ['build', str(self.spec), '-o', str(self.output), '--template', str(template), '--selfcheck', str(template)],
        ):
            with self.subTest(args=args):
                self.assertEqual(run(*args)[0], 1)
                self.assertEqual(template.read_bytes(), original)
                self.assertFalse(self.output.exists())

    def test_符号和硬链接不能绕过来源保护(self):
        symlink = self.root / 'symlink.json'; symlink.symlink_to(self.spec)
        hardlink = self.root / 'hardlink.json'; os.link(self.spec, hardlink)
        for alias in (symlink, hardlink):
            with self.subTest(alias=alias):
                self.assertEqual(run('report', str(self.spec), '-o', str(alias))[0], 1)
                self.assertEqual(self.spec.read_bytes(), self.original)
                self.assertEqual(alias.read_bytes(), self.original)

    def test_默认报告名也不能覆盖来源(self):
        source = Path(str(self.output) + '.selfcheck.md'); source.write_bytes(self.original)
        self.assertEqual(run('build', str(source), '-o', str(self.output))[0], 1)
        self.assertEqual(source.read_bytes(), self.original)
        self.assertFalse(self.output.exists())

    def test_正常重新生成已有输出仍可用(self):
        self.output.write_bytes(b'old output')
        self.assertEqual(run('build', str(self.spec), '-o', str(self.output))[0], 0)
        wb = openpyxl.load_workbook(self.output)
        self.assertEqual(wb.active['H2'].value, 'E'); wb.close()
        self.assertEqual(self.spec.read_bytes(), self.original)

    def test_模板旧项目名不能覆盖当前规格(self):
        template = self.root / 'template.xlsx'
        wb = openpyxl.Workbook()
        wb.active.append(['一级模块', '功能过程', '数据移动类型'])
        wb.active.append(['旧项目', '旧过程', 'E'])
        wb.save(template); wb.close()
        for extra, expected in (([], cosmic.STARTER_SPEC['project']),
                                (['--project', '当前项目'], '当前项目'),
                                (['--l1', '指定业务域'], '指定业务域')):
            with self.subTest(extra=extra):
                self.assertEqual(run('build', str(self.spec), '-o', str(self.output),
                                     '--template', str(template), *extra)[0], 0)
                wb = openpyxl.load_workbook(self.output)
                self.assertEqual(wb.active['A2'].value, expected)
                wb.close()
