"""Excel 导出必须保留业务原文，无法无损写出时保护已有产物。"""
import copy
import sys
import tempfile
import unittest
from pathlib import Path

SCRIPTS = Path(__file__).resolve().parents[2] / "skills/cosmic-split/scripts"
sys.path.insert(0, str(SCRIPTS))

import openpyxl
import builder
import cosmic
import sheet
import spec


class ExportIntegrityTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="cosmic-export-integrity-")
        self.path = Path(self.tmp.name) / "result.xlsx"
        self.rows = builder.build(spec.load_from_dict(copy.deepcopy(cosmic.STARTER_SPEC))).rows

    def tearDown(self):
        self.tmp.cleanup()

    def write(self, project="项目"):
        return sheet.write(self.rows, self.path, sheet.Template(l1_value=project), trace=True)

    def test_公式和错误字样在所有业务列保留为文本(self):
        for literal in ("=1+1", "#N/A", "#REF!", "=HYPERLINK(\"https://example.com\")"):
            with self.subTest(literal=literal):
                for row in self.rows:
                    for key in ("l2", "l3", "user", "event", "func", "desc", "group", "attrs", *sheet.TRACE_LABELS):
                        setattr(row, key, literal)
                self.write(project=literal)
                for data_only in (False, True):
                    workbook = openpyxl.load_workbook(self.path, data_only=data_only)
                    try:
                        worksheet = workbook.active
                        for column in (1, 2, 3, 4, 5, 6, 7, 9, 10, 13, 14, 15, 16, 17):
                            cell = worksheet.cell(2, column)
                            self.assertEqual((cell.value, cell.data_type), (literal, "s"))
                        self.assertEqual((worksheet["L2"].value, worksheet["L2"].data_type), (4, "n"))
                    finally:
                        workbook.close()

    def test_中英文和换行制表符原样保留(self):
        value = "中文 English 😀\n第二行\t字段\r\n第三行"
        self.rows[0].desc = value
        self.rows[0].attrs = "中文字段,EnglishID"
        self.write()
        workbook = openpyxl.load_workbook(self.path, data_only=True)
        try:
            self.assertEqual(workbook.active["G2"].value, value)
            self.assertEqual(workbook.active["J2"].value, "中文字段,EnglishID")
        finally:
            workbook.close()

    def test_32767字符可以无损写出(self):
        value = "长" * 32767
        self.rows[0].desc = value
        self.write()
        workbook = openpyxl.load_workbook(self.path)
        try:
            self.assertEqual(workbook.active["G2"].value, value)
        finally:
            workbook.close()

    def test_超长文本和拼接属性报错且不截断已有产物(self):
        self.write()
        original = self.path.read_bytes()
        for field, value in (("desc", "长" * 32768), ("attrs", ",".join(["字段" * 8192] * 2))):
            with self.subTest(field=field):
                previous = getattr(self.rows[-1], field)
                setattr(self.rows[-1], field, value)
                with self.assertRaises(SystemExit) as caught:
                    self.write()
                self.assertIn("FP-001", str(caught.exception))
                self.assertIn(field, str(caught.exception))
                self.assertIn("32767", str(caught.exception))
                self.assertEqual(self.path.read_bytes(), original)
                setattr(self.rows[-1], field, previous)

    def test_不支持的控制字符和XML字符不创建输出(self):
        for code in (0x00, 0x08, 0x0B, 0x0C, 0x0E, 0x1F, 0xD800, 0xDFFF, 0xFFFE, 0xFFFF):
            with self.subTest(code=code):
                self.rows[-1].source = "需求" + chr(code) + "原文"
                with self.assertRaises(SystemExit) as caught:
                    self.write()
                message = str(caught.exception)
                self.assertIn("FP-001", message)
                self.assertIn("source", message)
                self.assertIn(f"U+{code:04X}", message)
                self.assertFalse(self.path.exists())

    def test_XML不支持的字符不会损坏已有输出(self):
        self.write()
        original = self.path.read_bytes()
        self.rows[-1].desc = "非法\uffff原文"
        with self.assertRaises(SystemExit):
            self.write()
        self.assertEqual(self.path.read_bytes(), original)

    def test_项目名也需要预检(self):
        with self.assertRaises(SystemExit) as caught:
            self.write(project="项目\x00名称")
        self.assertIn("l1", str(caught.exception))
        self.assertFalse(self.path.exists())

    def heights(self, template=None):
        sheet.write(self.rows, self.path, template or sheet.Template(l1_value="项目"), trace=True)
        workbook = openpyxl.load_workbook(self.path)
        try:
            return [workbook.active.row_dimensions[row].height for row in range(2, 6)]
        finally:
            workbook.close()

    def test_短内容维持紧凑行高(self):
        self.assertEqual(self.heights(), [20, 20, 20, 20])

    def test_中文长描述与显式换行来源增加所在行高度(self):
        self.rows[0].desc = "接收业务用户提交的地址信息并核对详细内容" * 3
        self.rows[1].source = "需求说明第1段\n需求说明第2段\n需求说明第3段"
        template = sheet.Template(l1_value="项目")
        template.widths['source'] = 12
        heights = self.heights(template)
        self.assertGreater(heights[0], 50)
        self.assertGreater(heights[1], 45)
        self.assertEqual(heights[2:], [20, 20])
        workbook = openpyxl.load_workbook(self.path)
        try:
            self.assertEqual(workbook.active.column_dimensions['O'].width, 12)
        finally:
            workbook.close()

    def test_列宽增加减少换行且中文比同长度ASCII占更宽(self):
        self.rows[0].desc = "字" * 30
        narrow = sheet.Template(l1_value="项目")
        narrow.widths['desc'] = 20
        narrow_height = self.heights(narrow)[0]
        wide = sheet.Template(l1_value="项目")
        wide.widths['desc'] = 80
        self.assertLess(self.heights(wide)[0], narrow_height)
        self.rows[0].desc = "a" * 30
        self.assertLess(self.heights(narrow)[0], narrow_height)

    def test_来源中英文混排给按词换行预留一行(self):
        self.rows[0].source = "本任务提供的需求 R1/R2；R1：发送订单资料（客户编码、送货地址）"
        self.rows[1].source = "需求第一行\n需求第二行"
        heights = self.heights()
        self.assertGreaterEqual(heights[0], 48)
        self.assertLess(heights[1], 35)
        self.assertEqual(heights[2:], [20, 20])

    def test_纵向合并文本利用整个区域高度(self):
        for row in self.rows:
            row.event = "触发事件\n第二行\n第三行\n第四行"
        self.assertEqual(self.heights(), [20, 20, 20, 20])
        for row in self.rows:
            row.event = "\n".join(["触发事件"] * 10)
        heights = self.heights()
        self.assertTrue(all(20 < height < 50 for height in heights))
        self.assertGreater(sum(heights), 150)

    def test_超长内容行高不超过Excel上限且保留原文(self):
        self.rows[0].desc = "原文" * 1000
        for row in self.rows:
            row.event = "\n".join(["触发"] * 1000)
        heights = self.heights()
        self.assertEqual(heights, [409, 409, 409, 409])
        workbook = openpyxl.load_workbook(self.path)
        try:
            self.assertEqual(workbook.active['G2'].value, self.rows[0].desc)
            self.assertEqual(workbook.active['E2'].value, self.rows[0].event)
        finally:
            workbook.close()

    def test_inspect提示使用实际的一级模块参数(self):
        self.assertIn("--l1", sheet.describe(sheet.Template()))
        self.assertNotIn("--project", sheet.describe(sheet.Template()))


if __name__ == "__main__":
    unittest.main()
