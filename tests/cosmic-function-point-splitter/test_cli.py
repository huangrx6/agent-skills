"""CLI 行为与退出码。

跑法：
    python3 -m unittest discover -s tests/cosmic-function-point-splitter -v

钉的是「承诺的行为」：校验不通过就不落盘、init 不覆盖已有文件、audit 能读回自己写出的表。
"""

from __future__ import annotations

import contextlib
import io
import json
import os
import shutil
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.join(os.path.dirname(os.path.dirname(HERE)), "skills", os.path.basename(HERE))
sys.path.insert(0, os.path.join(SKILL, "scripts"))

import cosmic  # noqa: E402
import sheet  # noqa: E402


def run(*argv: str) -> tuple[int, str, str]:
    out, err = io.StringIO(), io.StringIO()
    with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
        code = cosmic.main(list(argv))
    return code, out.getvalue(), err.getvalue()


class CliTestCase(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp(prefix="cosmic-cli-")

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)

    def _path(self, name: str) -> str:
        return os.path.join(self.dir, name)

    def _starter(self) -> str:
        # 同一个测试里可能拿两次，重复 init 会因为「已存在」而失败，所以先看文件在不在
        spec = self._path("spec.json")
        if not os.path.isfile(spec):
            code, _out, _err = run("init", spec)
            self.assertEqual(code, 0)
        return spec

    def _broken(self) -> str:
        spec = self._starter()
        with open(spec, encoding="utf8") as handle:
            payload = json.load(handle)
        payload["processes"][0]["movements"][0]["type"] = "R"
        path = self._path("broken.json")
        with open(path, "w", encoding="utf8") as handle:
            handle.write(json.dumps(payload, ensure_ascii=False))
        return path


class TestInitAndCheck(CliTestCase):
    def test_init_写出骨架(self):
        spec = self._starter()
        self.assertTrue(os.path.isfile(spec))
        with open(spec, encoding="utf8") as handle:
            payload = json.load(handle)
        self.assertIn("measurement", payload)
        self.assertIn("processes", payload)

    def test_init_不覆盖已有文件(self):
        spec = self._starter()
        code, _out, err = run("init", spec)
        self.assertEqual(code, 1)
        self.assertIn("已存在", err)

    def test_init_force_覆盖(self):
        spec = self._starter()
        code, _out, _err = run("init", spec, "--force")
        self.assertEqual(code, 0)

    def test_check_骨架通过并给出功能点数(self):
        code, out, _err = run("check", self._starter())
        self.assertEqual(code, 0)
        # check 打的是按二级模块的分布，不是三级模块清单
        self.assertIn("示例一级模块", out)
        self.assertIn("4 CFP", out)
        self.assertIn("校验通过", out)

    def test_check_报出缺少触发输入(self):
        code, out, _err = run("check", self._broken(), "--max-issues", "50")
        self.assertEqual(code, 1)
        self.assertIn("未以 E 起", out)

    def test_规格引用错误时直接终止(self):
        spec = self._starter()
        with open(spec, encoding="utf8") as handle:
            payload = json.load(handle)
        payload["processes"][0]["l3"] = "不存在的模块"
        path = self._path("bad.json")
        with open(path, "w", encoding="utf8") as handle:
            handle.write(json.dumps(payload, ensure_ascii=False))
        with self.assertRaises(SystemExit) as ctx:
            run("check", path)
        self.assertIn("不存在的模块", str(ctx.exception))


class TestBuild(CliTestCase):
    def test_build_写出文件并报告核实状态(self):
        output = self._path("out.xlsx")
        code, out, _err = run("build", self._starter(), "-o", output)
        self.assertEqual(code, 0)
        self.assertTrue(os.path.isfile(output))
        self.assertIn("1 个功能过程", out)
        self.assertIn("4 行", out)
        self.assertIn("4 CFP", out)
        self.assertIn("已确认标记 0 CFP；待核实草稿 4 CFP", out)

    def test_校验不通过就不落盘(self):
        output = self._path("out.xlsx")
        code, _out, err = run("build", self._broken(), "-o", output)
        self.assertEqual(code, 1)
        self.assertFalse(os.path.isfile(output), "校验失败时不应写出文件")
        self.assertIn("未写出文件", err)
        self.assertIn("未以 E 起", err)

    def test_force_带病写出(self):
        output = self._path("out.xlsx")
        code, out, _err = run("build", self._broken(), "-o", output, "--force")
        self.assertEqual(code, 0)
        self.assertTrue(os.path.isfile(output))
        self.assertIn("跳过校验", out)

    def test_selfcheck_写出自检清单(self):
        output = self._path("out.xlsx")
        report = self._path("自检清单.md")
        code, _out, _err = run("build", self._starter(), "-o", output, "--selfcheck", report)
        self.assertEqual(code, 0)
        with open(report, encoding="utf8") as handle:
            text = handle.read()
        self.assertIn("自检清单", text)
        self.assertIn("逐项来源与计量依据", text)
        self.assertIn("待核实 4 CFP", text)

    def test_一级模块取值来自规格(self):
        output = self._path("out.xlsx")
        self.assertEqual(run("build", self._starter(), "-o", output)[0], 0)
        import openpyxl

        worksheet = openpyxl.load_workbook(output).active
        self.assertEqual(worksheet.cell(row=2, column=1).value, "示例项目")

    def test_report_缺省打到标准输出_给_o_时写文件(self):
        code, out, _err = run("report", self._starter())
        self.assertEqual(code, 0)
        self.assertIn("自检清单", out)

        report = self._path("自检清单.md")
        code, _out, _err = run("report", self._starter(), "-o", report)
        self.assertEqual(code, 0)
        self.assertTrue(os.path.isfile(report))

    def test_report_用_o_而不是_selfcheck(self):
        # report 只有一个写出入口，避免两个名字都能写文件却只生效一个
        with self.assertRaises(SystemExit):
            run("report", self._starter(), "--selfcheck", self._path("x.md"))


class TestAudit(CliTestCase):
    def test_audit_读回自己写出的表无问题(self):
        output = self._path("out.xlsx")
        run("build", self._starter(), "-o", output)
        code, out, _err = run("audit", output)
        self.assertEqual(code, 0, out)
        self.assertIn("未发现结构问题", out)

    def test_audit_抓出缺失描述(self):
        import openpyxl

        output = self._path("out.xlsx")
        run("build", self._starter(), "-o", output)
        workbook = openpyxl.load_workbook(output)
        worksheet = workbook.active
        # 删除第一行描述，审计必须报告缺失。
        worksheet.cell(row=2, column=7, value="")
        workbook.save(output)
        code, out, _err = run("audit", output, "--max-issues", "50")
        self.assertEqual(code, 1)
        self.assertIn("缺少功能过程、描述或数据组", out)

    def test_audit_读不到数据行时报错(self):
        import openpyxl

        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.cell(row=1, column=1, value="一级模块")
        path = self._path("empty.xlsx")
        workbook.save(path)
        code, _out, err = run("audit", path)
        self.assertEqual(code, 1)
        self.assertIn("没读到数据行", err)


class TestInspectCli(CliTestCase):
    def test_inspect_打印列布局(self):
        import openpyxl

        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        for column, label in enumerate(
            ["一级模块", "二级模块", "三级模块", "功能用户", "触发事件", "功能过程",
             "子过程描述", "数据移动类型", "数据组", "数据属性", "复用度", "CFP"], start=1
        ):
            worksheet.cell(row=1, column=column, value=label)
        path = self._path("t.xlsx")
        workbook.save(path)
        code, out, _err = run("inspect", path)
        self.assertEqual(code, 0)
        self.assertIn("逻辑字段", out)
        self.assertIn("表头行：1", out)

    def test_build_用模板布局(self):
        import openpyxl

        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        for column, label in enumerate(
            ["二级模块", "三级模块", "功能过程", "标识"], start=1
        ):
            worksheet.cell(row=1, column=column, value=label)
        worksheet.cell(row=2, column=1, value="模块")
        worksheet.cell(row=2, column=4, value="E")
        template = self._path("t.xlsx")
        workbook.save(template)

        output = self._path("out.xlsx")
        code, _out, _err = run(
            "build", self._starter(), "-o", output, "--template", template
        )
        self.assertEqual(code, 0)
        result = openpyxl.load_workbook(output).active
        headers = [result.cell(row=1, column=c).value for c in range(1, result.max_column + 1)]
        # 前 4 列沿用模板（「标识」被认成数据移动类型），缺的必填列接在后面
        self.assertEqual(headers[:4], ["二级模块", "三级模块", "功能过程", "数据移动类型"])
        self.assertIn("CFP", headers, "缺的必填列应被追加")
        self.assertEqual(headers[4:], ["一级模块", "功能用户", "触发事件", "子过程描述",
                                   "数据组", "数据属性", "CFP", "过程编号", "核实状态", "需求依据", "兴趣对象", "重复移动理由",
                                   "操作类型", "操作依据", "接口证据"])
        # 复用度是常量列，模板没有就不追加
        self.assertNotIn("复用度", headers)


if __name__ == "__main__":
    unittest.main()
