"""已有表审计必须保留缺损证据，且不能跨过程补造元数据。"""

import copy
import sys
import tempfile
import unittest
from pathlib import Path

SCRIPTS = Path(__file__).resolve().parents[2] / 'skills/cosmic-function-point-splitter/scripts'
sys.path.insert(0, str(SCRIPTS))

import builder
import checks
import cosmic
import openpyxl
import sheet
import spec


class AuditIntegrityTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix='cosmic-audit-integrity-')
        self.addCleanup(self.tmp.cleanup)
        self.path = Path(self.tmp.name) / 'audit.xlsx'

    def write(self, second=False):
        value = copy.deepcopy(cosmic.STARTER_SPEC)
        if second:
            process = copy.deepcopy(value['processes'][0])
            process.update(id='FP-002', name='地址修改', trigger='业务用户提交地址修改')
            value['processes'].append(process)
        sheet.write(builder.build(spec.load_from_dict(value)).rows, self.path,
                    sheet.Template(l1_value='项目'))

    def edit(self, change):
        workbook = openpyxl.load_workbook(self.path)
        change(workbook.active)
        workbook.save(self.path)
        workbook.close()

    @staticmethod
    def unmerge_columns(worksheet, columns):
        for merged in list(worksheet.merged_cells.ranges):
            if merged.min_col in columns:
                worksheet.unmerge_cells(str(merged))

    def read(self):
        rows = cosmic.read_existing_table(self.path)
        return rows, checks.check(rows)

    def test_有过程名或编号的缺损行不得丢弃(self):
        for key in ('func', 'process_id'):
            with self.subTest(key=key):
                self.write()
                layout = sheet.inspect(self.path).layout

                def append(worksheet):
                    worksheet.cell(6, layout[key], '遗漏业务' if key == 'func' else 'FP-002')
                    worksheet.cell(6, layout['cfp'], 999)

                self.edit(append)
                rows, issues = self.read()
                self.assertEqual(len(rows), 5)
                self.assertEqual(rows[-1].declared_cfp, 999)
                self.assertTrue(any('非空行 第 6 行' in issue for issue in issues), issues)
                self.assertTrue(any('非法数据移动类型 第 6 行' in issue for issue in issues), issues)

    def test_只有CFP的非空行也需核实(self):
        self.write()
        self.edit(lambda worksheet: setattr(worksheet['L6'], 'value', 999))
        rows, issues = self.read()
        self.assertEqual(len(rows), 5)
        self.assertTrue(any('缺损过程或汇总行' in issue for issue in issues), issues)

    def test_不能根据合计文字静默忽略行(self):
        self.write()

        def append(worksheet):
            worksheet['F6'] = '合计'
            worksheet['L6'] = 999

        self.edit(append)
        rows, issues = self.read()
        self.assertEqual(rows[-1].func, '合计')
        self.assertTrue(any('非空行 第 6 行' in issue for issue in issues), issues)

    def test_新编号不能继承上一过程的名称用户和触发(self):
        self.write(second=True)

        def clear(worksheet):
            self.unmerge_columns(worksheet, (4, 5, 6))
            for row in range(6, 10):
                for column in (4, 5, 6):
                    worksheet.cell(row, column).value = None

        self.edit(clear)
        rows, issues = self.read()
        second = rows[4]
        self.assertEqual(second.process_id, 'FP-002')
        self.assertEqual((second.func, second.user, second.event), ('', '', ''))
        self.assertTrue(any('缺少功能过程' in issue for issue in issues), issues)
        self.assertTrue(any('功能用户、触发事件' in issue for issue in issues), issues)

    def test_无编号的新过程名也不能继承用户触发(self):
        self.write(second=True)

        def clear(worksheet):
            self.unmerge_columns(worksheet, (4, 5))
            for row in range(2, 10):
                worksheet.cell(row, 13).value = None
            for row in range(6, 10):
                worksheet.cell(row, 4).value = None
                worksheet.cell(row, 5).value = None

        self.edit(clear)
        rows, issues = self.read()
        self.assertEqual(rows[4].func, '地址修改')
        self.assertEqual((rows[4].user, rows[4].event), ('', ''))
        self.assertTrue(any('功能用户、触发事件' in issue for issue in issues), issues)

    def test_同过程连续行仍可继承省略的元数据(self):
        self.write()

        def clear(worksheet):
            self.unmerge_columns(worksheet, (2, 3, 4, 5, 6))
            for row in range(3, 6):
                worksheet.cell(row, 13).value = None

        self.edit(clear)
        rows, issues = self.read()
        for row in rows:
            self.assertEqual((row.process_id, row.func, row.user, row.event),
                             ('FP-001', '地址新增', '业务用户', '业务用户提交新地址'))
        self.assertEqual(issues, [])

    def test_实际合并单元格可跨过程保留相同功能用户(self):
        self.write(second=True)
        rows, issues = self.read()
        self.assertEqual(rows[4].user, '业务用户')
        self.assertEqual(rows[4].event, '业务用户提交地址修改')
        self.assertEqual(rows[4].func, '地址修改')
        self.assertEqual(issues, [])

    def test_核心列缺失需明确报告(self):
        for key in ('l2', 'l3', 'user', 'event', 'object'):
            with self.subTest(key=key):
                self.write()
                column = sheet.inspect(self.path).layout[key]

                def clear(worksheet):
                    self.unmerge_columns(worksheet, (column,))
                    for row in range(1, 6):
                        worksheet.cell(row, column).value = None

                self.edit(clear)
                _, issues = self.read()
                self.assertTrue(any('原表缺少核心列' in issue and sheet.HEADER_LABELS[key] in issue
                                    for issue in issues), issues)

    def test_核心列存在但值为空也不能通过(self):
        for key in ('l2', 'l3', 'user', 'event', 'object'):
            with self.subTest(key=key):
                self.write()
                column = sheet.inspect(self.path).layout[key]

                def clear(worksheet):
                    self.unmerge_columns(worksheet, (column,))
                    for row in range(2, 6):
                        worksheet.cell(row, column).value = None

                self.edit(clear)
                _, issues = self.read()
                self.assertTrue(any('缺少核心字段' in issue and sheet.HEADER_LABELS[key] in issue
                                    for issue in issues), issues)


if __name__ == '__main__':
    unittest.main()
