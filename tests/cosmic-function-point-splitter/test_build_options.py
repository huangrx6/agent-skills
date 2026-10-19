"""生成选项：排序、功能用户写法、确认门与严格模式。"""

from __future__ import annotations

import contextlib
import copy
import io
import json
import os
import shutil
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.path.join(os.path.dirname(os.path.dirname(HERE)), "skills", os.path.basename(HERE))
sys.path.insert(0, os.path.join(SKILL, "scripts"))

# 静态分析器看不到 sys.path 挂载的同级模块。
# pyright: reportMissingImports=false

import builder  # noqa: E402
import checks  # noqa: E402
import cosmic  # noqa: E402
import openpyxl  # noqa: E402
import sheet  # noqa: E402
import spec  # noqa: E402


def active(workbook):
    """openpyxl 的类型里 workbook.active 可能为空；测试里直接取默认表。"""
    worksheet = workbook.active
    if worksheet is None:  # pragma: no cover - 由 openpyxl 保证
        raise AssertionError("工作簿没有默认工作表")
    return worksheet


def run(*argv: str) -> tuple[int, str, str]:
    out, err = io.StringIO(), io.StringIO()
    with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
        try:
            code = cosmic.main(list(argv))
        except SystemExit as exc:
            code = exc.code if isinstance(exc.code, int) else 1
    return code, out.getvalue(), err.getvalue()


def two_processes(query_first: bool = True):
    value = copy.deepcopy(cosmic.STARTER_SPEC)
    query = copy.deepcopy(value['processes'][0])
    query.update(id='FP-002', name='地址列表查询', op='queryList',
                 trigger='业务用户打开地址列表')
    query['movements'] = [
        {"type": "E", "object": "地址", "group": "地址资料",
         "attributes": ["地址编码"], "description": "输入地址查询请求"},
        {"type": "R", "object": "地址", "group": "地址资料",
         "attributes": ["地址编码"], "description": "读取地址信息"},
        {"type": "X", "object": "地址", "group": "地址资料",
         "attributes": ["地址编码"], "description": "返回地址查询结果"},
    ]
    if query_first:
        value['processes'] = [query, value['processes'][0]]
    else:
        value['processes'] = [value['processes'][0], query]
    return value


class BuildOptionsCase(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp(prefix="cosmic-options-")
        self.addCleanup(shutil.rmtree, self.dir, True)

    def write_spec(self, value) -> str:
        path = os.path.join(self.dir, "spec.json")
        with open(path, "w", encoding="utf8") as handle:
            json.dump(value, handle, ensure_ascii=False)
        return path

    def output(self, name: str = "out.xlsx") -> str:
        return os.path.join(self.dir, name)


class TestOrder(BuildOptionsCase):
    def test_默认保留声明顺序并提示送审惯例偏离(self):
        value = two_processes(query_first=True)
        rows = builder.build(spec.load_from_dict(value)).rows
        heads = [row.func for row in rows if row.is_process_head]
        self.assertEqual(heads, ['地址列表查询', '地址新增'])
        hints = checks.hints(rows)
        self.assertTrue(any('不符合送审惯例' in hint for hint in hints), hints)

    def test_crud排序只移动过程块且不改计量(self):
        value = two_processes(query_first=True)
        spec_value = spec.load_from_dict(value)
        before = builder.build(spec_value)
        after = builder.build(spec_value, order="crud")
        self.assertEqual([row.func for row in after.rows if row.is_process_head],
                         ['地址新增', '地址列表查询'])
        self.assertEqual(after.total_cfp, before.total_cfp)
        self.assertEqual(len(after.rows), len(before.rows))
        self.assertEqual(sorted(row.move for row in after.rows),
                         sorted(row.move for row in before.rows))
        self.assertFalse(any('不符合送审惯例' in hint for hint in checks.hints(after.rows)))

    def test_命令行crud排序写出(self):
        path = self.write_spec(two_processes(query_first=True))
        output = self.output()
        code, out, _err = run("build", path, "-o", output, "--order", "crud")
        self.assertEqual(code, 0, out)
        rows = cosmic.read_existing_table(output)
        self.assertEqual([row.func for row in rows if row.is_process_head],
                         ['地址新增', '地址列表查询'])


class TestUserFormat(BuildOptionsCase):
    def test_双行功能用户写入发起者与接收者(self):
        path = self.write_spec(cosmic.STARTER_SPEC)
        output = self.output()
        code, out, _err = run("build", path, "-o", output, "--user-format", "dual")
        self.assertEqual(code, 0, out)
        workbook = openpyxl.load_workbook(output)
        worksheet = active(workbook)
        column = sheet.inspect(output).layout['user']
        self.assertEqual(worksheet.cell(row=2, column=column).value,
                         "发起者：业务用户\n接收者：地址管理")
        workbook.close()
        self.assertEqual(run("audit", output)[0], 0)

    def test_未知写法被拒绝(self):
        path = self.write_spec(cosmic.STARTER_SPEC)
        code, _out, err = run("build", path, "-o", self.output(), "--user-format", "triple")
        self.assertEqual(code, 2)
        self.assertIn("invalid choice", err)


class TestGates(BuildOptionsCase):
    def test_require_confirmed阻止未核实过程(self):
        path = self.write_spec(cosmic.STARTER_SPEC)
        output = self.output()
        code, _out, err = run("build", path, "-o", output, "--require-confirmed")
        self.assertEqual(code, 1)
        self.assertFalse(os.path.exists(output))
        self.assertIn("未标 confirmed", err)

    def test_全部确认后可以写出(self):
        value = copy.deepcopy(cosmic.STARTER_SPEC)
        value['processes'][0].update(status='confirmed', source='需求 2.1 地址新增')
        path = self.write_spec(value)
        output = self.output()
        code, out, _err = run("build", path, "-o", output, "--require-confirmed")
        self.assertEqual(code, 0, out)
        self.assertTrue(os.path.exists(output))

    def test_strict把提示当失败且不写表(self):
        value = copy.deepcopy(cosmic.STARTER_SPEC)
        value['processes'][0].pop('op_note')
        path = self.write_spec(value)
        output = self.output()
        code, _out, err = run("build", path, "-o", output, "--strict")
        self.assertEqual(code, 1)
        self.assertFalse(os.path.exists(output))
        self.assertIn("提示", err)

        code, out, _err = run("build", path, "-o", output)
        self.assertEqual(code, 0, out)
        self.assertTrue(os.path.exists(output))

    def test_check严格模式退出码跟随提示(self):
        value = copy.deepcopy(cosmic.STARTER_SPEC)
        value['processes'][0].pop('op_note')
        path = self.write_spec(value)
        self.assertEqual(run("check", path)[0], 0)
        self.assertEqual(run("check", path, "--strict")[0], 1)

    def test_audit严格模式也接受提示(self):
        path = self.write_spec(cosmic.STARTER_SPEC)
        output = self.output()
        self.assertEqual(run("build", path, "-o", output)[0], 0)
        self.assertEqual(run("audit", output, "--strict")[0], 0, "骨架示例没有提示")

    def test_操作依据及证据往返保持且严格检查一致(self):
        value = copy.deepcopy(cosmic.STARTER_SPEC)
        value['processes'][0]['evidence'] = ['api-address-create']
        path = self.write_spec(value)
        output = self.output()
        self.assertEqual(run('build', path, '-o', output)[0], 0)
        row = cosmic.read_existing_table(output)[0]
        self.assertEqual(row.op, 'insert')
        self.assertEqual(row.op_note, value['processes'][0]['op_note'])
        self.assertEqual(row.evidence, 'api-address-create')
        value['processes'][0].pop('op_note')
        self.write_spec(value)
        self.assertEqual(run('build', path, '-o', output)[0], 0)
        self.assertEqual(run('check', path, '--strict')[0], 1)
        self.assertEqual(run('audit', output, '--strict')[0], 1)
        workbook = openpyxl.load_workbook(output)
        active(workbook).cell(2, sheet.inspect(output).layout['op'], '未知标签')
        workbook.save(output); workbook.close()
        self.assertEqual(run('audit', output, '--strict')[0], 1)


if __name__ == "__main__":
    unittest.main()
